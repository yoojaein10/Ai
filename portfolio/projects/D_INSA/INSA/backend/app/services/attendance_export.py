from calendar import monthrange
from collections import defaultdict
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Optional

import holidays
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.db.models import Department, Employee
from app.services.attendance import (
    _half_day_multiplier,
    _hhmmss,
    _minutes_between,
)


_LATE_THRESHOLD_HHMMSS = "090000"
_MORNING_ARRIVAL_THRESHOLD_HHMMSS = "093000"
_CHECKOUT_THRESHOLD_HHMMSS = "180000"
_KO_WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]
_THIN = Side(style="thin", color="999999")
_BORDER_ALL = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_FILL = PatternFill("solid", fgColor="1a1a2e")
_GROUP_FILL = PatternFill("solid", fgColor="e6e6f0")
_WEEKEND_FILL = PatternFill("solid", fgColor="fff1f5")
_HOLIDAY_FILL = PatternFill("solid", fgColor="fff1f5")
_LEAVE_FILL = PatternFill("solid", fgColor="e6f4ff")
_TRIP_FILL = PatternFill("solid", fgColor="fff7e6")


def _is_late(first_time: Optional[str]) -> bool:
    if not first_time or len(first_time) < 6:
        return False
    return first_time[:6] > _LATE_THRESHOLD_HHMMSS


def _kr_holiday_set(year: int, month: int) -> set[date]:
    kr = holidays.KR(years=[year])
    last = monthrange(year, month)[1]
    return {
        d
        for d in kr
        if d.year == year and d.month == month and 1 <= d.day <= last
    }


def _business_days_in_month(year: int, month: int) -> int:
    hset = _kr_holiday_set(year, month)
    last = monthrange(year, month)[1]
    count = 0
    for i in range(last):
        cur = date(year, month, 1) + timedelta(days=i)
        if cur.weekday() < 5 and cur not in hset:
            count += 1
    return count


def _classify_schedule(gubun: str, bigo: Optional[str]) -> str:
    """SCHEDULE.gubun + Bigo → 집계 카테고리."""
    g = (gubun or "").strip()
    b = (bigo or "").replace(" ", "")
    if g == "휴가":
        return "반차" if "반차" in b else "연차"
    if g == "공가":
        return "공가"
    if g in ("특가", "NPL휴가"):
        return "경조휴가" if "경조" in b else "특가"
    if g == "가족돌봄":
        return "가족돌봄"
    return ""


def _collect_month_data(
    att_db: Session,
    apw_db: Session,
    insa_db: Session,
    year: int,
    month: int,
):
    last_day = monthrange(year, month)[1]
    start_date = date(year, month, 1)
    end_date = date(year, month, last_day)
    next_month = start_date + timedelta(days=last_day)
    start_str = start_date.strftime("%Y%m%d")
    end_str = end_date.strftime("%Y%m%d")

    per: dict[str, dict] = defaultdict(
        lambda: {
            "emp_no": None,
            "tenter_by_date": {},
            "sched_by_date": {},
            "trip_by_date": defaultdict(list),
            "late_days": 0,
            "missing_auth_days": 0,
            "work_day_set": set(),
            "work_minutes": 0,
            "counts": defaultdict(float),
            "footer_notes": defaultdict(set),
        }
    )

    tenter_rows = att_db.execute(
        text(
            """
            SELECT e_idno,
                   MAX(e_name) AS e_name,
                   e_date,
                   MIN(e_time) AS first_t,
                   MAX(e_time) AS last_t,
                   COUNT(*) AS tag_count
            FROM tenter
            WHERE e_date BETWEEN :s AND :e
              AND e_id <> -1
              AND e_idno IS NOT NULL
              AND e_idno <> ''
            GROUP BY e_idno, e_date
            """
        ),
        {"s": start_str, "e": end_str},
    ).mappings().all()

    for r in tenter_rows:
        name = (r["e_name"] or "").strip()
        if not name:
            continue
        entry = per[name]
        if entry["emp_no"] is None:
            entry["emp_no"] = r["e_idno"]
        d_str = r["e_date"]
        try:
            d_obj = datetime.strptime(d_str, "%Y%m%d").date()
        except ValueError:
            continue
        entry["tenter_by_date"][d_obj] = {
            "first": r["first_t"],
            "last": r["last_t"],
            "tag_count": int(r["tag_count"]),
        }
        entry["work_day_set"].add(d_obj)
        entry["work_minutes"] += _minutes_between(r["first_t"], r["last_t"])
        if _is_late(r["first_t"]):
            entry["late_days"] += 1
        if int(r["tag_count"]) <= 1:
            entry["missing_auth_days"] += 1

    sched_rows = apw_db.execute(
        text(
            """
            SELECT name, [date], gubun, Bigo
            FROM APW_IW_SCHEDULE
            WHERE [date] >= :s AND [date] < :ne
            """
        ),
        {"s": start_date, "ne": next_month},
    ).mappings().all()

    for r in sched_rows:
        name = (r["name"] or "").strip()
        if not name:
            continue
        d_val = r["date"]
        d_key = d_val.date() if isinstance(d_val, datetime) else d_val
        gubun = (r["gubun"] or "").strip()
        bigo = (r["Bigo"] or "").strip()
        entry = per[name]
        entry["sched_by_date"][d_key] = {"gubun": gubun, "bigo": bigo}

        cat = _classify_schedule(gubun, bigo)
        if cat:
            mult = _half_day_multiplier(bigo) if cat in ("연차", "공가") else 1.0
            if cat == "반차":
                entry["counts"]["반차"] += 1
            else:
                entry["counts"][cat] += mult

        bigo_nospace = bigo.replace(" ", "")
        note_map = [
            ("육아휴직", "육아휴직"),
            ("재택근무", "재택근무"),
            ("단축근무", "단축근무"),
        ]
        for keyword, note_key in note_map:
            if keyword in bigo_nospace or keyword in gubun:
                entry["footer_notes"][note_key].add(d_key)

    trip_rows = apw_db.execute(
        text(
            """
            SELECT Name, CAST(CulDate AS date) AS d, Gubun, Address
            FROM APW_IW_DAYCULJANG
            WHERE CulDate >= :s AND CulDate < :ne
            ORDER BY CulDate
            """
        ),
        {"s": start_date, "ne": next_month},
    ).mappings().all()

    for r in trip_rows:
        name = (r["Name"] or "").strip()
        if not name:
            continue
        addr = (r["Address"] or "").strip()
        gubun = r["Gubun"] if r["Gubun"] is not None else 0
        per[name]["trip_by_date"][r["d"]].append((addr, int(gubun)))

    names = list(per.keys())
    dept_info: dict[str, tuple[Optional[str], Optional[str]]] = {}
    if names:
        matched = (
            insa_db.query(Employee.emp_no, Employee.name_ko, Department.name)
            .outerjoin(Department, Employee.dept_id == Department.id)
            .filter(Employee.name_ko.in_(names))
            .all()
        )
        for emp_no, name_ko, dept_name in matched:
            dept_info[name_ko] = (emp_no, dept_name)

    return {
        "per": per,
        "dept_info": dept_info,
        "start_date": start_date,
        "last_day": last_day,
        "business_days": _business_days_in_month(year, month),
        "holiday_set": _kr_holiday_set(year, month),
    }


def _group_ranges(days: set[date]) -> list[tuple[date, date]]:
    """연속 일자를 구간으로 압축."""
    if not days:
        return []
    sorted_days = sorted(days)
    ranges: list[tuple[date, date]] = []
    start = prev = sorted_days[0]
    for d in sorted_days[1:]:
        if (d - prev).days == 1:
            prev = d
            continue
        ranges.append((start, prev))
        start = prev = d
    ranges.append((start, prev))
    return ranges


def _fmt_range(rng: tuple[date, date]) -> str:
    s, e = rng
    if s == e:
        return s.strftime("%m/%d")
    return f"{s.strftime('%m/%d')}~{e.strftime('%m/%d')}"


def build_summary_xlsx(
    att_db: Session,
    apw_db: Session,
    insa_db: Session,
    year: int,
    month: int,
) -> bytes:
    data = _collect_month_data(att_db, apw_db, insa_db, year, month)
    per = data["per"]
    dept_info = data["dept_info"]
    business_days = data["business_days"]

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d} 출퇴근현황"

    headers = [
        "No.",
        "소속",
        "성명",
        "출근",
        "지각",
        "연차",
        "반차",
        "공가",
        "특가",
        "경조휴가",
        "가족돌봄",
        "출근률",
        "인증누락",
    ]
    title = f"{year}년 {month}월 출퇴근 현황 (영업일 {business_days}일)"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    tcell = ws.cell(row=1, column=1, value=title)
    tcell.font = Font(size=14, bold=True)
    tcell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = Font(color="ffffff", bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER_ALL

    rows_sorted = []
    for name, e in per.items():
        emp_no, dept_name = dept_info.get(name, (e["emp_no"], None))
        rows_sorted.append(
            {
                "name": name,
                "emp_no": emp_no or e["emp_no"] or "",
                "dept": dept_name or "(미매칭)",
                "work_days": len(e["work_day_set"]),
                "late": e["late_days"],
                "leave": round(e["counts"]["연차"], 1),
                "half": int(e["counts"]["반차"]),
                "gong": round(e["counts"]["공가"], 1),
                "special": round(e["counts"]["특가"], 1),
                "bereavement": round(e["counts"]["경조휴가"], 1),
                "family": round(e["counts"]["가족돌봄"], 1),
                "missing": e["missing_auth_days"],
            }
        )
    rows_sorted.sort(key=lambda r: (r["dept"], r["name"]))

    row_idx = 3
    current_dept: Optional[str] = None
    seq = 0
    for r in rows_sorted:
        if r["dept"] != current_dept:
            current_dept = r["dept"]
            ws.merge_cells(
                start_row=row_idx,
                start_column=1,
                end_row=row_idx,
                end_column=len(headers),
            )
            gcell = ws.cell(row=row_idx, column=1, value=f"▶ {current_dept}")
            gcell.font = Font(bold=True)
            gcell.fill = _GROUP_FILL
            gcell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            gcell.border = _BORDER_ALL
            row_idx += 1

        seq += 1
        attendance_rate = (
            f"{round(r['work_days'] / business_days * 100, 1)}%"
            if business_days > 0
            else "-"
        )
        values = [
            seq,
            r["dept"],
            r["name"],
            r["work_days"],
            r["late"],
            r["leave"],
            r["half"],
            r["gong"],
            r["special"],
            r["bereavement"],
            r["family"],
            attendance_rate,
            r["missing"],
        ]
        for col_idx, v in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=v)
            cell.border = _BORDER_ALL
            if col_idx in (1, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        row_idx += 1

    widths = [5, 20, 12, 8, 8, 8, 8, 8, 8, 10, 10, 10, 10]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D3"

    # 하단 주석: 육아휴직/재택근무/단축근무 기간
    row_idx += 1
    ws.merge_cells(
        start_row=row_idx,
        start_column=1,
        end_row=row_idx,
        end_column=len(headers),
    )
    hcell = ws.cell(row=row_idx, column=1, value="▣ 특이사항")
    hcell.font = Font(bold=True, size=11)
    hcell.fill = _GROUP_FILL
    hcell.alignment = Alignment(horizontal="left", indent=1)
    row_idx += 1

    note_labels = [
        ("육아휴직", "육아휴직"),
        ("재택근무", "재택근무"),
        ("단축근무", "단축근무"),
    ]
    any_note = False
    for note_key, label in note_labels:
        for r in rows_sorted:
            e = per[r["name"]]
            days = e["footer_notes"].get(note_key)
            if not days:
                continue
            any_note = True
            ranges = _group_ranges(days)
            range_text = ", ".join(_fmt_range(x) for x in ranges)
            text_val = f"{label}: {r['dept']} {r['name']} ({range_text})"
            ws.merge_cells(
                start_row=row_idx,
                start_column=1,
                end_row=row_idx,
                end_column=len(headers),
            )
            ncell = ws.cell(row=row_idx, column=1, value=text_val)
            ncell.alignment = Alignment(horizontal="left", indent=2)
            row_idx += 1

    if not any_note:
        ws.merge_cells(
            start_row=row_idx,
            start_column=1,
            end_row=row_idx,
            end_column=len(headers),
        )
        ws.cell(row=row_idx, column=1, value="해당 사항 없음").alignment = Alignment(
            horizontal="left", indent=2
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _trip_label(trips: list[tuple[str, int]], gubun_filter: set[int]) -> Optional[str]:
    """trips 중 주어진 Gubun에 해당하는 주소 라벨."""
    if not trips:
        return None
    filtered = [a for a, g in trips if g in gubun_filter and a]
    if filtered:
        return ", ".join(filtered)
    # 주소가 비어있는 출장도 존재할 수 있음
    if any(g in gubun_filter for _, g in trips):
        return "출장"
    return None


def _history_cell_value(
    cur: date,
    kind: str,
    tenter: Optional[dict],
    sched: Optional[dict],
    trips: list[tuple[str, int]],
    is_weekend: bool,
    is_holiday: bool,
) -> tuple[str, Optional[PatternFill]]:
    """한 셀에 들어갈 값과 배경색."""
    if is_weekend or is_holiday:
        return "-", _WEEKEND_FILL if is_weekend else _HOLIDAY_FILL

    if sched:
        gubun = sched["gubun"]
        bigo = sched["bigo"]
        bigo_ns = bigo.replace(" ", "")
        if gubun == "휴가":
            if "오전반차" in bigo_ns:
                label = "오전반차"
            elif "오후반차" in bigo_ns:
                label = "오후반차"
            elif "반차" in bigo_ns:
                label = "반차"
            else:
                label = bigo or "연차"
        elif gubun == "공가":
            label = bigo or "공가"
        elif gubun in ("특가", "NPL휴가"):
            label = bigo or gubun
        else:
            label = bigo or gubun or "-"
        return label, _LEAVE_FILL

    # Gubun 1=오전, 2=오후, 3=종일 → 출근 셀: 1/3, 퇴근 셀: 2/3
    am_label = _trip_label(trips, {1, 3})
    pm_label = _trip_label(trips, {2, 3})

    if kind == "출근":
        first = tenter["first"] if tenter else None
        if first and first[:6] <= _MORNING_ARRIVAL_THRESHOLD_HHMMSS:
            hhmmss = _hhmmss(first)
            return (hhmmss[:5] if hhmmss else "-"), None
        # 09:30까지 출근 태그가 없거나 그 이후 태그 → 오전 출장 있으면 대체
        if am_label:
            return am_label, _TRIP_FILL
        if first:
            hhmmss = _hhmmss(first)
            return (hhmmss[:5] if hhmmss else "-"), None
        return "-", None

    # kind == "퇴근"
    if tenter:
        last = tenter["last"]
        first = tenter["first"]
        if last and last != first:
            is_real_checkout = last[:6] >= _CHECKOUT_THRESHOLD_HHMMSS
            if not is_real_checkout and pm_label:
                return pm_label, _TRIP_FILL
            hhmmss = _hhmmss(last)
            return (hhmmss[:5] if hhmmss else "-"), None
        # 태그 1건뿐 → 오후 출장 있으면 대체
        if pm_label:
            return pm_label, _TRIP_FILL
        return "-", None

    if pm_label:
        return pm_label, _TRIP_FILL
    return "-", None


def build_history_xlsx(
    att_db: Session,
    apw_db: Session,
    insa_db: Session,
    year: int,
    month: int,
) -> bytes:
    data = _collect_month_data(att_db, apw_db, insa_db, year, month)
    per = data["per"]
    dept_info = data["dept_info"]
    start_date = data["start_date"]
    last_day = data["last_day"]
    holiday_set = data["holiday_set"]

    wb = Workbook()
    ws = wb.active
    ws.title = f"{year}-{month:02d} 출퇴근이력"

    fixed_headers = ["부서", "이름", "구분"]
    day_headers = [str(i) for i in range(1, last_day + 1)]
    total_cols = len(fixed_headers) + last_day

    title = f"{year}년 {month}월 출퇴근 이력"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    tcell = ws.cell(row=1, column=1, value=title)
    tcell.font = Font(size=14, bold=True)
    tcell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for col_idx, h in enumerate(fixed_headers + day_headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = Font(color="ffffff", bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER_ALL

    weekday_row = 3
    for col_idx, d_num in enumerate(range(1, last_day + 1), start=len(fixed_headers) + 1):
        cur = date(year, month, d_num)
        wday = _KO_WEEKDAYS[cur.weekday()]
        cell = ws.cell(row=weekday_row, column=col_idx, value=wday)
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER_ALL
        if cur.weekday() >= 5 or cur in holiday_set:
            cell.fill = _WEEKEND_FILL
            cell.font = Font(color="c41d7f", bold=True)
    for col_idx in range(1, len(fixed_headers) + 1):
        ws.cell(row=weekday_row, column=col_idx).border = _BORDER_ALL

    rows_sorted = []
    for name, e in per.items():
        emp_no, dept_name = dept_info.get(name, (e["emp_no"], None))
        rows_sorted.append(
            {
                "name": name,
                "dept": dept_name or "(미매칭)",
                "entry": e,
            }
        )
    rows_sorted.sort(key=lambda r: (r["dept"], r["name"]))

    row_idx = 4
    for r in rows_sorted:
        e = r["entry"]
        ws.merge_cells(
            start_row=row_idx, start_column=1, end_row=row_idx + 1, end_column=1
        )
        ws.merge_cells(
            start_row=row_idx, start_column=2, end_row=row_idx + 1, end_column=2
        )
        dcell = ws.cell(row=row_idx, column=1, value=r["dept"])
        ncell = ws.cell(row=row_idx, column=2, value=r["name"])
        for c in (dcell, ncell):
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = _BORDER_ALL

        for kind_offset, kind in enumerate(["출근", "퇴근"]):
            rr = row_idx + kind_offset
            kcell = ws.cell(row=rr, column=3, value=kind)
            kcell.alignment = Alignment(horizontal="center", vertical="center")
            kcell.border = _BORDER_ALL

            for d_num in range(1, last_day + 1):
                cur = date(year, month, d_num)
                is_weekend = cur.weekday() >= 5
                is_holiday = cur in holiday_set
                tenter = e["tenter_by_date"].get(cur)
                sched = e["sched_by_date"].get(cur)
                trips = e["trip_by_date"].get(cur, [])

                val, fill = _history_cell_value(
                    cur, kind, tenter, sched, trips, is_weekend, is_holiday
                )
                col_idx = len(fixed_headers) + d_num
                cell = ws.cell(row=rr, column=col_idx, value=val)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = _BORDER_ALL
                if fill:
                    cell.fill = fill
                if is_weekend or is_holiday:
                    cell.font = Font(color="c41d7f")

        row_idx += 2

    ws.column_dimensions[get_column_letter(1)].width = 16
    ws.column_dimensions[get_column_letter(2)].width = 10
    ws.column_dimensions[get_column_letter(3)].width = 7
    for i in range(len(fixed_headers) + 1, total_cols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 8
    ws.freeze_panes = "D4"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
