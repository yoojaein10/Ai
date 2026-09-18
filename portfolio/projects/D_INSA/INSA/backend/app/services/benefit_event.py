from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy import bindparam, or_, text
from sqlalchemy.orm import Session

from app.db.models import BenefitEvent, BenefitItem, Department, Employee
from app.schemas.benefit import (
    BenefitEventCreate,
    BenefitEventOut,
    BenefitEventUpdate,
    LinkedSchedule,
    UploadResult,
)


_SCHEDULE_GUBUNS = ("특가", "NPL휴가")


def _to_out(
    ev: BenefitEvent,
    emp: Optional[Employee],
    dept_name: Optional[str],
    item: Optional[BenefitItem],
    linked: list[LinkedSchedule],
) -> BenefitEventOut:
    return BenefitEventOut(
        id=ev.id,
        employee_id=ev.employee_id,
        emp_no=emp.emp_no if emp else None,
        emp_name=emp.name_ko if emp else None,
        dept_name=dept_name,
        item_id=ev.item_id,
        item_code=item.code if item else None,
        item_name=item.name if item else None,
        event_type=ev.event_type,
        target_person=ev.target_person,
        event_date=ev.event_date,
        amount=ev.amount,
        leave_days=ev.leave_days,
        remark=ev.remark,
        linked_schedules=linked,
    )


def _find_linked_schedules(
    apw_db: Session, name: str, event_date: date, window_days: int = 7
) -> list[LinkedSchedule]:
    if not name:
        return []
    start = event_date - timedelta(days=window_days)
    end = event_date + timedelta(days=window_days + 1)
    stmt = text(
        """
        SELECT [date], gubun, Bigo
        FROM APW_IW_SCHEDULE
        WHERE name = :name
          AND [date] >= :s AND [date] < :e
          AND gubun IN :gubuns
        ORDER BY [date]
        """
    ).bindparams(bindparam("gubuns", expanding=True))
    rows = apw_db.execute(
        stmt,
        {"name": name, "s": start, "e": end, "gubuns": list(_SCHEDULE_GUBUNS)},
    ).mappings().all()
    result: list[LinkedSchedule] = []
    for r in rows:
        bigo = (r["Bigo"] or "").replace(" ", "")
        if "경조" not in bigo:
            continue
        d_val = r["date"]
        d_key = d_val.date() if isinstance(d_val, datetime) else d_val
        result.append(
            LinkedSchedule(
                schedule_date=d_key,
                gubun=(r["gubun"] or "").strip(),
                bigo=r["Bigo"],
            )
        )
    return result


def list_events(
    db: Session,
    apw_db: Session,
    year: Optional[int] = None,
    employee_id: Optional[int] = None,
    event_type: Optional[str] = None,
    keyword: Optional[str] = None,
) -> list[BenefitEventOut]:
    query = (
        db.query(BenefitEvent, Employee, Department, BenefitItem)
        .outerjoin(Employee, BenefitEvent.employee_id == Employee.id)
        .outerjoin(Department, Employee.dept_id == Department.id)
        .outerjoin(BenefitItem, BenefitEvent.item_id == BenefitItem.id)
    )
    if year:
        query = query.filter(
            BenefitEvent.event_date >= date(year, 1, 1),
            BenefitEvent.event_date <= date(year, 12, 31),
        )
    if employee_id:
        query = query.filter(BenefitEvent.employee_id == employee_id)
    if event_type:
        query = query.filter(BenefitEvent.event_type == event_type)
    if keyword:
        like = f"%{keyword}%"
        query = query.filter(
            or_(
                Employee.name_ko.ilike(like),
                Employee.emp_no.ilike(like),
                BenefitEvent.target_person.ilike(like),
            )
        )

    rows = query.order_by(BenefitEvent.event_date.desc(), BenefitEvent.id.desc()).all()
    return [
        _to_out(
            ev,
            emp,
            dept.name if dept else None,
            item,
            _find_linked_schedules(apw_db, emp.name_ko if emp else "", ev.event_date),
        )
        for ev, emp, dept, item in rows
    ]


def get_event(db: Session, event_id: int) -> Optional[BenefitEvent]:
    return db.query(BenefitEvent).filter(BenefitEvent.id == event_id).first()


def create_event(db: Session, data: BenefitEventCreate) -> BenefitEvent:
    ev = BenefitEvent(**data.model_dump())
    db.add(ev)
    db.commit()
    db.refresh(ev)
    return ev


def update_event(
    db: Session, event_id: int, data: BenefitEventUpdate
) -> Optional[BenefitEvent]:
    ev = get_event(db, event_id)
    if ev is None:
        return None
    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(ev, field, value)
    db.commit()
    db.refresh(ev)
    return ev


def delete_event(db: Session, event_id: int) -> bool:
    ev = get_event(db, event_id)
    if ev is None:
        return False
    db.delete(ev)
    db.commit()
    return True


# ── Excel Upload ─────────────────────────────────────────
_HEADERS = ["사번", "이름", "경조구분", "대상자", "일자", "금액", "휴가일수", "비고"]


def _parse_date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(v) -> Optional[Decimal]:
    if v is None or v == "":
        return None
    try:
        s = str(v).replace(",", "").strip()
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def upload_events(db: Session, content: bytes) -> UploadResult:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    errors: list[str] = []
    created = 0
    skipped = 0

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return UploadResult(created=0, skipped=0, errors=["빈 파일입니다"])

    header_row = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {h: header_row.index(h) for h in _HEADERS if h in header_row}
    for required in ["이름", "경조구분", "일자"]:
        if required not in idx:
            return UploadResult(
                created=0, skipped=0, errors=[f"필수 헤더 누락: {required}"]
            )

    for i, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or c == "" for c in row):
            continue

        def get(col: str):
            if col not in idx:
                return None
            return row[idx[col]]

        emp_no = (str(get("사번")).strip() if get("사번") else "") or None
        name = (str(get("이름")).strip() if get("이름") else "") or None
        event_type = (
            str(get("경조구분")).strip() if get("경조구분") else ""
        ) or None
        event_date = _parse_date(get("일자"))

        if not name or not event_type or not event_date:
            skipped += 1
            errors.append(f"{i}행: 이름/경조구분/일자 누락")
            continue

        emp_query = db.query(Employee)
        if emp_no:
            emp_query = emp_query.filter(Employee.emp_no == emp_no)
        else:
            emp_query = emp_query.filter(Employee.name_ko == name)
        emp = emp_query.first()
        if emp is None:
            skipped += 1
            errors.append(f"{i}행: 직원 매칭 실패 ({emp_no or ''} {name})")
            continue

        item = (
            db.query(BenefitItem)
            .filter(
                BenefitItem.category == "경조",
                BenefitItem.event_type == event_type,
                BenefitItem.is_active == True,  # noqa: E712
            )
            .first()
        )

        ev = BenefitEvent(
            employee_id=emp.id,
            item_id=item.id if item else None,
            event_type=event_type,
            target_person=(
                str(get("대상자")).strip() if get("대상자") else None
            )
            or None,
            event_date=event_date,
            amount=_parse_decimal(get("금액")),
            leave_days=_parse_decimal(get("휴가일수")),
            remark=(str(get("비고")).strip() if get("비고") else None) or None,
        )
        db.add(ev)
        created += 1

    db.commit()
    return UploadResult(created=created, skipped=skipped, errors=errors[:50])
