from datetime import date, datetime
from io import BytesIO
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.db.models import BenefitHealth, Department, Employee
from app.schemas.benefit import (
    BenefitHealthCreate,
    BenefitHealthOut,
    BenefitHealthUpdate,
    UploadResult,
)


_ABNORMAL_RESULTS = {"요관찰", "유소견", "질환의심"}


def _to_out(
    rec: BenefitHealth,
    emp: Optional[Employee],
    dept_name: Optional[str],
) -> BenefitHealthOut:
    return BenefitHealthOut(
        id=rec.id,
        employee_id=rec.employee_id,
        emp_no=emp.emp_no if emp else None,
        emp_name=emp.name_ko if emp else None,
        dept_name=dept_name,
        check_year=rec.check_year,
        check_date=rec.check_date,
        provider=rec.provider,
        check_type=rec.check_type,
        result=rec.result,
        recheck_required=rec.recheck_required,
        recheck_date=rec.recheck_date,
        remark=rec.remark,
    )


def list_health(
    db: Session,
    year: Optional[int] = None,
    employee_id: Optional[int] = None,
    result: Optional[str] = None,
    recheck_only: bool = False,
    keyword: Optional[str] = None,
) -> list[BenefitHealthOut]:
    query = (
        db.query(BenefitHealth, Employee, Department)
        .outerjoin(Employee, BenefitHealth.employee_id == Employee.id)
        .outerjoin(Department, Employee.dept_id == Department.id)
    )
    if year:
        query = query.filter(BenefitHealth.check_year == year)
    if employee_id:
        query = query.filter(BenefitHealth.employee_id == employee_id)
    if result:
        query = query.filter(BenefitHealth.result == result)
    if recheck_only:
        query = query.filter(BenefitHealth.recheck_required == True)  # noqa: E712
    if keyword:
        like = f"%{keyword}%"
        query = query.filter(
            or_(Employee.name_ko.ilike(like), Employee.emp_no.ilike(like))
        )

    rows = query.order_by(
        BenefitHealth.check_year.desc(), BenefitHealth.check_date.desc()
    ).all()
    return [_to_out(r, emp, dept.name if dept else None) for r, emp, dept in rows]


def get_health(db: Session, rec_id: int) -> Optional[BenefitHealth]:
    return db.query(BenefitHealth).filter(BenefitHealth.id == rec_id).first()


def create_health(db: Session, data: BenefitHealthCreate) -> BenefitHealth:
    rec = BenefitHealth(**data.model_dump())
    db.add(rec)
    db.commit()
    db.refresh(rec)
    return rec


def update_health(
    db: Session, rec_id: int, data: BenefitHealthUpdate
) -> Optional[BenefitHealth]:
    rec = get_health(db, rec_id)
    if rec is None:
        return None
    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(rec, field, value)
    db.commit()
    db.refresh(rec)
    return rec


def delete_health(db: Session, rec_id: int) -> bool:
    rec = get_health(db, rec_id)
    if rec is None:
        return False
    db.delete(rec)
    db.commit()
    return True


# ── Excel Upload ─────────────────────────────────────────
_HEADERS = [
    "사번",
    "이름",
    "검진년도",
    "검진일",
    "검진기관",
    "검진종류",
    "결과",
    "재검필요",
    "재검일",
    "비고",
]


def _parse_date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_bool(v) -> bool:
    if v is None or v == "":
        return False
    s = str(v).strip().lower()
    return s in ("y", "yes", "true", "1", "재검", "필요", "o")


def upload_health(db: Session, content: bytes) -> UploadResult:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    errors: list[str] = []
    created = 0
    skipped = 0

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return UploadResult(created=0, skipped=0, errors=["빈 파일입니다"])

    header_row = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {h: header_row.index(h) for h in _HEADERS if h in header_row}
    for required in ["이름", "검진년도"]:
        if required not in idx:
            return UploadResult(
                created=0, skipped=0, errors=[f"필수 헤더 누락: {required}"]
            )

    for i, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or c == "" for c in row):
            continue

        def get(col: str):
            if col not in idx:
                return None
            return row[idx[col]]

        emp_no = (str(get("사번")).strip() if get("사번") else "") or None
        name = (str(get("이름")).strip() if get("이름") else "") or None
        year_raw = get("검진년도")
        try:
            check_year = int(str(year_raw).strip()) if year_raw is not None else None
        except ValueError:
            check_year = None

        if not name or not check_year:
            skipped += 1
            errors.append(f"{i}행: 이름/검진년도 누락 또는 형식 오류")
            continue

        emp_query = db.query(Employee)
        if emp_no:
            emp_query = emp_query.filter(Employee.emp_no == emp_no)
        else:
            emp_query = emp_query.filter(Employee.name_ko == name)
        emp = emp_query.first()
        if emp is None:
            skipped += 1
            errors.append(f"{i}행: 직원 매칭 실패 ({emp_no or ''} {name})")
            continue

        result_val = (str(get("결과")).strip() if get("결과") else None) or None
        recheck_input = get("재검필요")
        recheck = _parse_bool(recheck_input)
        if not recheck and result_val in _ABNORMAL_RESULTS:
            recheck = True

        rec = BenefitHealth(
            employee_id=emp.id,
            check_year=check_year,
            check_date=_parse_date(get("검진일")),
            provider=(str(get("검진기관")).strip() if get("검진기관") else None)
            or None,
            check_type=(
                str(get("검진종류")).strip() if get("검진종류") else None
            )
            or None,
            result=result_val,
            recheck_required=recheck,
            recheck_date=_parse_date(get("재검일")),
            remark=(str(get("비고")).strip() if get("비고") else None) or None,
        )
        db.add(rec)
        created += 1

    db.commit()
    return UploadResult(created=created, skipped=skipped, errors=errors[:50])
