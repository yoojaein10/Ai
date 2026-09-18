from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Optional

from openpyxl import load_workbook
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.db.models import Department, EduCourse, EduRecord, Employee
from app.schemas.edu import (
    EduRecordCreate,
    EduRecordOut,
    EduRecordUpdate,
    EduRecordUploadResult,
)


def _to_out(
    rec: EduRecord,
    emp: Optional[Employee] = None,
    dept_name: Optional[str] = None,
    course: Optional[EduCourse] = None,
) -> EduRecordOut:
    return EduRecordOut(
        id=rec.id,
        employee_id=rec.employee_id,
        emp_no=emp.emp_no if emp else None,
        emp_name=emp.name_ko if emp else None,
        dept_name=dept_name,
        course_id=rec.course_id,
        course_code=course.course_code if course else None,
        course_name_snapshot=rec.course_name_snapshot,
        category=course.category if course else None,
        start_date=rec.start_date,
        end_date=rec.end_date,
        hours=rec.hours,
        score=rec.score,
        result=rec.result,
        certificate_no=rec.certificate_no,
        remark=rec.remark,
    )


def list_records(
    db: Session,
    year: Optional[int] = None,
    employee_id: Optional[int] = None,
    course_id: Optional[int] = None,
    category: Optional[str] = None,
    keyword: Optional[str] = None,
) -> list[EduRecordOut]:
    query = (
        db.query(EduRecord, Employee, Department, EduCourse)
        .outerjoin(Employee, EduRecord.employee_id == Employee.id)
        .outerjoin(Department, Employee.dept_id == Department.id)
        .outerjoin(EduCourse, EduRecord.course_id == EduCourse.id)
    )
    if year:
        start = date(year, 1, 1)
        end = date(year, 12, 31)
        query = query.filter(
            or_(
                EduRecord.end_date.between(start, end),
                EduRecord.start_date.between(start, end),
            )
        )
    if employee_id:
        query = query.filter(EduRecord.employee_id == employee_id)
    if course_id:
        query = query.filter(EduRecord.course_id == course_id)
    if category:
        query = query.filter(EduCourse.category == category)
    if keyword:
        like = f"%{keyword}%"
        query = query.filter(
            or_(
                Employee.name_ko.ilike(like),
                Employee.emp_no.ilike(like),
                EduRecord.course_name_snapshot.ilike(like),
            )
        )

    rows = query.order_by(EduRecord.end_date.desc(), EduRecord.id.desc()).all()
    return [
        _to_out(rec, emp, dept.name if dept else None, course)
        for rec, emp, dept, course in rows
    ]


def get_record(db: Session, record_id: int) -> Optional[EduRecord]:
    return db.query(EduRecord).filter(EduRecord.id == record_id).first()


def create_record(db: Session, data: EduRecordCreate) -> EduRecord:
    payload = data.model_dump()
    if payload.get("course_id") and not payload.get("course_name_snapshot"):
        course = db.query(EduCourse).filter(EduCourse.id == payload["course_id"]).first()
        if course:
            payload["course_name_snapshot"] = course.course_name
    rec = EduRecord(**payload)
    db.add(rec)
    db.commit()
    db.refresh(rec)
    return rec


def update_record(
    db: Session, record_id: int, data: EduRecordUpdate
) -> Optional[EduRecord]:
    rec = get_record(db, record_id)
    if rec is None:
        return None
    for field, value in data.model_dump(exclude_unset=True).items():
        setattr(rec, field, value)
    db.commit()
    db.refresh(rec)
    return rec


def delete_record(db: Session, record_id: int) -> bool:
    rec = get_record(db, record_id)
    if rec is None:
        return False
    db.delete(rec)
    db.commit()
    return True


# ── Excel upload ──────────────────────────────────────────
_EXPECTED_HEADERS = [
    "사번",
    "이름",
    "과정코드",
    "과정명",
    "시작일",
    "종료일",
    "이수시간",
    "점수",
    "결과",
    "수료증번호",
    "비고",
]


def _parse_date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_decimal(v) -> Optional[Decimal]:
    if v is None or v == "":
        return None
    try:
        return Decimal(str(v).strip())
    except (InvalidOperation, ValueError):
        return None


def upload_records(db: Session, content: bytes) -> EduRecordUploadResult:
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    errors: list[str] = []
    created = 0
    skipped = 0

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return EduRecordUploadResult(created=0, skipped=0, errors=["빈 파일입니다"])

    header_row = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {h: header_row.index(h) for h in _EXPECTED_HEADERS if h in header_row}
    required = ["이름", "과정명"]
    missing = [h for h in required if h not in idx]
    if missing:
        return EduRecordUploadResult(
            created=0,
            skipped=0,
            errors=[f"필수 헤더 누락: {', '.join(missing)}"],
        )

    for i, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or c == "" for c in row):
            continue

        def get(col: str):
            if col not in idx:
                return None
            return row[idx[col]]

        emp_no = (str(get("사번")).strip() if get("사번") else "") or None
        name = (str(get("이름")).strip() if get("이름") else "") or None
        course_code = (str(get("과정코드")).strip() if get("과정코드") else "") or None
        course_name = (str(get("과정명")).strip() if get("과정명") else "") or None

        if not name or not course_name:
            skipped += 1
            errors.append(f"{i}행: 이름/과정명 누락")
            continue

        emp_query = db.query(Employee)
        if emp_no:
            emp_query = emp_query.filter(Employee.emp_no == emp_no)
        else:
            emp_query = emp_query.filter(Employee.name_ko == name)
        emp = emp_query.first()
        if emp is None:
            skipped += 1
            errors.append(f"{i}행: 직원 매칭 실패 ({emp_no or ''} {name})")
            continue

        course_id: Optional[int] = None
        if course_code:
            course = (
                db.query(EduCourse).filter(EduCourse.course_code == course_code).first()
            )
            if course:
                course_id = course.id

        rec = EduRecord(
            employee_id=emp.id,
            course_id=course_id,
            course_name_snapshot=course_name,
            start_date=_parse_date(get("시작일")),
            end_date=_parse_date(get("종료일")),
            hours=_parse_decimal(get("이수시간")),
            score=_parse_decimal(get("점수")),
            result=(str(get("결과")).strip() if get("결과") else None) or None,
            certificate_no=(
                str(get("수료증번호")).strip() if get("수료증번호") else None
            )
            or None,
            remark=(str(get("비고")).strip() if get("비고") else None) or None,
        )
        db.add(rec)
        created += 1

    db.commit()
    return EduRecordUploadResult(created=created, skipped=skipped, errors=errors[:50])
