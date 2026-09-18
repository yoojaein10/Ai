from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.db.models import BenefitHealth, Department, Employee
from app.schemas.benefit import BenefitHealthCreate, BenefitHealthUpdate
from app.services import benefit_health as svc


@pytest.fixture
def dept(db_session):
    d = Department(code="D-H", name="개발팀")
    db_session.add(d)
    db_session.commit()
    return d


@pytest.fixture
def emp(db_session, dept):
    e = Employee(
        emp_no="H001", name_ko="검진자", hire_date=date(2020, 1, 1),
        dept_id=dept.id,
    )
    db_session.add(e)
    db_session.commit()
    return e


@pytest.fixture
def emp2(db_session, dept):
    e = Employee(
        emp_no="H002", name_ko="다른사람", hire_date=date(2021, 1, 1),
        dept_id=dept.id,
    )
    db_session.add(e)
    db_session.commit()
    return e


def _make_record(db, emp_id, year=2025, result=None, recheck=False):
    rec = svc.create_health(db, BenefitHealthCreate(
        employee_id=emp_id, check_year=year,
        check_date=date(year, 6, 15), provider="A병원",
        check_type="일반", result=result, recheck_required=recheck,
    ))
    return rec


# ── parsers ───────────────────────────────────────────────


def test_parse_date_handles_none_and_empty():
    assert svc._parse_date(None) is None
    assert svc._parse_date("") is None


def test_parse_date_handles_datetime_and_date():
    assert svc._parse_date(datetime(2025, 6, 15, 10, 0)) == date(2025, 6, 15)
    assert svc._parse_date(date(2025, 6, 15)) == date(2025, 6, 15)


def test_parse_date_string_formats():
    assert svc._parse_date("2025-06-15") == date(2025, 6, 15)
    assert svc._parse_date("2025/06/15") == date(2025, 6, 15)
    assert svc._parse_date("2025.06.15") == date(2025, 6, 15)
    assert svc._parse_date("20250615") == date(2025, 6, 15)
    assert svc._parse_date("not-a-date") is None


def test_parse_bool_truthy_values():
    assert svc._parse_bool("Y") is True
    assert svc._parse_bool("yes") is True
    assert svc._parse_bool("재검") is True
    assert svc._parse_bool("필요") is True
    assert svc._parse_bool("1") is True


def test_parse_bool_falsy_values():
    assert svc._parse_bool(None) is False
    assert svc._parse_bool("") is False
    assert svc._parse_bool("N") is False
    assert svc._parse_bool("false") is False


# ── CRUD ──────────────────────────────────────────────────


def test_list_filters(db_session, emp, emp2):
    _make_record(db_session, emp.id, 2024, result="정상A")
    _make_record(db_session, emp.id, 2025, result="요관찰", recheck=True)
    _make_record(db_session, emp2.id, 2025, result="정상B")
    assert len(svc.list_health(db_session)) == 3
    assert len(svc.list_health(db_session, year=2025)) == 2
    assert len(svc.list_health(db_session, employee_id=emp.id)) == 2
    assert len(svc.list_health(db_session, result="요관찰")) == 1
    assert len(svc.list_health(db_session, recheck_only=True)) == 1


def test_list_keyword_matches_name_and_emp_no(db_session, emp, emp2):
    _make_record(db_session, emp.id)
    _make_record(db_session, emp2.id)
    by_name = svc.list_health(db_session, keyword="검진자")
    by_no = svc.list_health(db_session, keyword="H002")
    assert len(by_name) == 1
    assert by_name[0].emp_name == "검진자"
    assert len(by_no) == 1
    assert by_no[0].emp_no == "H002"


def test_get_health(db_session, emp):
    rec = _make_record(db_session, emp.id)
    assert svc.get_health(db_session, rec.id).id == rec.id
    assert svc.get_health(db_session, 99999) is None


def test_create_and_update(db_session, emp):
    rec = _make_record(db_session, emp.id, result="요관찰")
    updated = svc.update_health(
        db_session, rec.id,
        BenefitHealthUpdate(result="정상B", recheck_required=False),
    )
    assert updated.result == "정상B"
    assert updated.recheck_required is False


def test_update_not_found(db_session):
    assert svc.update_health(db_session, 99999, BenefitHealthUpdate(result="x")) is None


def test_delete(db_session, emp):
    rec = _make_record(db_session, emp.id)
    assert svc.delete_health(db_session, rec.id) is True
    assert svc.get_health(db_session, rec.id) is None


def test_delete_not_found(db_session):
    assert svc.delete_health(db_session, 99999) is False


# ── Excel upload ──────────────────────────────────────────


def _build_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


_HEADER = ["사번", "이름", "검진년도", "검진일", "검진기관", "검진종류",
           "결과", "재검필요", "재검일", "비고"]


def test_upload_empty_file(db_session):
    wb = Workbook()
    buf = BytesIO()
    wb.active.delete_rows(1, wb.active.max_row + 1)
    wb.save(buf)
    res = svc.upload_health(db_session, buf.getvalue())
    assert res.created == 0
    assert "빈 파일" in res.errors[0]


def test_upload_missing_required_header(db_session):
    content = _build_xlsx([
        ["사번", "검진년도"],  # "이름" missing
        ["H001", 2025],
    ])
    res = svc.upload_health(db_session, content)
    assert res.created == 0
    assert "필수 헤더 누락" in res.errors[0]


def test_upload_skips_blank_rows(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["H001", "검진자", 2025, "2025-06-15", "A병원", "일반", "정상A", "N", "", ""],
        [None, None, None, None, None, None, None, None, None, None],
    ])
    res = svc.upload_health(db_session, content)
    assert res.created == 1
    assert res.skipped == 0


def test_upload_skips_when_name_or_year_missing(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["H001", "", 2025, "", "", "", "", "", "", ""],
        ["H001", "검진자", "notayear", "", "", "", "", "", "", ""],
    ])
    res = svc.upload_health(db_session, content)
    assert res.created == 0
    assert res.skipped == 2


def test_upload_skips_when_employee_not_matched(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["UNKNOWN", "유령직원", 2025, "", "", "", "", "", "", ""],
    ])
    res = svc.upload_health(db_session, content)
    assert res.created == 0
    assert res.skipped == 1
    assert "직원 매칭 실패" in res.errors[0]


def test_upload_matches_by_name_when_no_emp_no(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["", "검진자", 2025, "2025-06-15", "B병원", "종합", "정상A", "N", "", ""],
    ])
    res = svc.upload_health(db_session, content)
    assert res.created == 1
    rows = svc.list_health(db_session, employee_id=emp.id)
    assert rows[0].provider == "B병원"


def test_upload_auto_recheck_for_abnormal_results(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["H001", "검진자", 2025, "2025-06-15", "A병원", "일반", "유소견", "N", "", ""],
    ])
    res = svc.upload_health(db_session, content)
    assert res.created == 1
    rec = svc.list_health(db_session, employee_id=emp.id)[0]
    assert rec.recheck_required is True


def test_upload_explicit_recheck_flag(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["H001", "검진자", 2025, "2025-06-15", "A병원", "일반", "정상A", "재검", "2025-09-15", "비고"],
    ])
    res = svc.upload_health(db_session, content)
    assert res.created == 1
    rec = svc.list_health(db_session, employee_id=emp.id)[0]
    assert rec.recheck_required is True
    assert rec.recheck_date == date(2025, 9, 15)
    assert rec.remark == "비고"
