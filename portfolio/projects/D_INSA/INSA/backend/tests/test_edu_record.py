from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.db.models import Department, EduCourse, EduRecord, Employee
from app.schemas.edu import EduRecordCreate, EduRecordUpdate
from app.services import edu_record as svc


@pytest.fixture
def dept(db_session):
    d = Department(code="D-EDU", name="개발팀")
    db_session.add(d)
    db_session.commit()
    return d


@pytest.fixture
def emp(db_session, dept):
    e = Employee(
        emp_no="EDU001", name_ko="이수자",
        hire_date=date(2020, 1, 1), dept_id=dept.id,
    )
    db_session.add(e)
    db_session.commit()
    return e


@pytest.fixture
def emp2(db_session, dept):
    e = Employee(
        emp_no="EDU002", name_ko="다른수강자",
        hire_date=date(2021, 1, 1), dept_id=dept.id,
    )
    db_session.add(e)
    db_session.commit()
    return e


@pytest.fixture
def course_law(db_session):
    c = EduCourse(
        course_code="C-LAW", course_name="법정필수교육",
        category="법정", hours=Decimal("8.0"),
    )
    db_session.add(c)
    db_session.commit()
    return c


@pytest.fixture
def course_job(db_session):
    c = EduCourse(
        course_code="C-JOB", course_name="직무역량",
        category="직무", hours=Decimal("4.0"),
    )
    db_session.add(c)
    db_session.commit()
    return c


def _make(db, *, emp_id, course_id=None, course_name="과정", year=2025, month=6,
          hours="8.0", result="수료"):
    return svc.create_record(db, EduRecordCreate(
        employee_id=emp_id, course_id=course_id,
        course_name_snapshot=course_name,
        start_date=date(year, month, 1),
        end_date=date(year, month, 5),
        hours=Decimal(hours), result=result,
    ))


# ── parsers ───────────────────────────────────────────────


def test_parse_date_handles_none_and_empty():
    assert svc._parse_date(None) is None
    assert svc._parse_date("") is None


def test_parse_date_datetime_passthrough():
    assert svc._parse_date(datetime(2025, 6, 1)) == date(2025, 6, 1)
    assert svc._parse_date(date(2025, 6, 1)) == date(2025, 6, 1)


def test_parse_date_string_formats():
    for s, expected in [
        ("2025-06-01", date(2025, 6, 1)),
        ("2025/06/01", date(2025, 6, 1)),
        ("2025.06.01", date(2025, 6, 1)),
        ("20250601", date(2025, 6, 1)),
    ]:
        assert svc._parse_date(s) == expected
    assert svc._parse_date("garbage") is None


def test_parse_decimal_valid_and_invalid():
    assert svc._parse_decimal(None) is None
    assert svc._parse_decimal("") is None
    assert svc._parse_decimal("8.5") == Decimal("8.5")
    assert svc._parse_decimal(8) == Decimal("8")
    assert svc._parse_decimal("not-a-number") is None


# ── create / get / update / delete ────────────────────────


def test_create_uses_course_name_when_snapshot_empty(db_session, emp, course_law):
    rec = svc.create_record(db_session, EduRecordCreate(
        employee_id=emp.id, course_id=course_law.id,
        course_name_snapshot="",  # empty - should be filled from course
        hours=Decimal("8.0"),
    ))
    assert rec.course_name_snapshot == "법정필수교육"


def test_create_keeps_explicit_snapshot(db_session, emp, course_law):
    rec = _make(db_session, emp_id=emp.id, course_id=course_law.id,
                course_name="명시이름")
    assert rec.course_name_snapshot == "명시이름"


def test_get_record(db_session, emp, course_law):
    rec = _make(db_session, emp_id=emp.id, course_id=course_law.id)
    assert svc.get_record(db_session, rec.id).id == rec.id
    assert svc.get_record(db_session, 99999) is None


def test_update_partial(db_session, emp, course_law):
    rec = _make(db_session, emp_id=emp.id, course_id=course_law.id, result="진행중")
    updated = svc.update_record(db_session, rec.id, EduRecordUpdate(result="수료"))
    assert updated.result == "수료"


def test_update_not_found(db_session):
    assert svc.update_record(db_session, 99999, EduRecordUpdate(result="수료")) is None


def test_delete(db_session, emp, course_law):
    rec = _make(db_session, emp_id=emp.id, course_id=course_law.id)
    assert svc.delete_record(db_session, rec.id) is True
    assert svc.get_record(db_session, rec.id) is None


def test_delete_not_found(db_session):
    assert svc.delete_record(db_session, 99999) is False


# ── list filters ──────────────────────────────────────────


def test_list_filter_by_year(db_session, emp, course_law):
    _make(db_session, emp_id=emp.id, course_id=course_law.id, year=2024)
    _make(db_session, emp_id=emp.id, course_id=course_law.id, year=2025)
    assert len(svc.list_records(db_session, year=2025)) == 1
    assert len(svc.list_records(db_session)) == 2


def test_list_filter_by_employee(db_session, emp, emp2, course_law):
    _make(db_session, emp_id=emp.id, course_id=course_law.id)
    _make(db_session, emp_id=emp2.id, course_id=course_law.id)
    assert len(svc.list_records(db_session, employee_id=emp.id)) == 1


def test_list_filter_by_course_and_category(db_session, emp, course_law, course_job):
    _make(db_session, emp_id=emp.id, course_id=course_law.id)
    _make(db_session, emp_id=emp.id, course_id=course_job.id)
    by_course = svc.list_records(db_session, course_id=course_law.id)
    by_cat = svc.list_records(db_session, category="법정")
    assert len(by_course) == 1
    assert len(by_cat) == 1


def test_list_filter_by_keyword(db_session, emp, emp2, course_law):
    _make(db_session, emp_id=emp.id, course_id=course_law.id)
    _make(db_session, emp_id=emp2.id, course_id=course_law.id)
    by_name = svc.list_records(db_session, keyword="이수자")
    by_no = svc.list_records(db_session, keyword="EDU002")
    assert len(by_name) == 1
    assert by_name[0].emp_name == "이수자"
    assert len(by_no) == 1
    assert by_no[0].emp_no == "EDU002"


def test_list_returns_dept_and_category(db_session, emp, course_law):
    _make(db_session, emp_id=emp.id, course_id=course_law.id)
    rows = svc.list_records(db_session)
    assert rows[0].dept_name == "개발팀"
    assert rows[0].category == "법정"
    assert rows[0].course_code == "C-LAW"


# ── Excel upload ──────────────────────────────────────────


def _build_xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


_HEADER = ["사번", "이름", "과정코드", "과정명", "시작일", "종료일",
           "이수시간", "점수", "결과", "수료증번호", "비고"]


def test_upload_empty_file(db_session):
    wb = Workbook()
    wb.active.delete_rows(1, wb.active.max_row + 1)
    buf = BytesIO()
    wb.save(buf)
    res = svc.upload_records(db_session, buf.getvalue())
    assert res.created == 0
    assert res.errors == ["빈 파일입니다"]


def test_upload_missing_required_header(db_session):
    content = _build_xlsx([
        ["사번", "이름"],  # missing 과정명
        ["EDU001", "이수자"],
    ])
    res = svc.upload_records(db_session, content)
    assert res.created == 0
    assert "필수 헤더 누락" in res.errors[0]


def test_upload_skips_blank_rows(db_session, emp, course_law):
    content = _build_xlsx([
        _HEADER,
        ["EDU001", "이수자", "C-LAW", "법정필수교육",
         "2025-06-01", "2025-06-05", "8", "90", "수료", "CERT-1", ""],
        [None] * len(_HEADER),
    ])
    res = svc.upload_records(db_session, content)
    assert res.created == 1
    assert res.skipped == 0


def test_upload_skips_when_name_or_course_missing(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["EDU001", "", "", "과정", "", "", "", "", "", "", ""],
        ["EDU001", "이수자", "", "", "", "", "", "", "", "", ""],
    ])
    res = svc.upload_records(db_session, content)
    assert res.created == 0
    assert res.skipped == 2


def test_upload_skips_when_employee_not_matched(db_session):
    content = _build_xlsx([
        _HEADER,
        ["UNKNOWN", "유령", "", "과정X", "", "", "", "", "", "", ""],
    ])
    res = svc.upload_records(db_session, content)
    assert res.created == 0
    assert res.skipped == 1
    assert "직원 매칭 실패" in res.errors[0]


def test_upload_matches_by_name_when_no_emp_no(db_session, emp, course_law):
    content = _build_xlsx([
        _HEADER,
        ["", "이수자", "C-LAW", "법정필수교육",
         "2025-06-01", "2025-06-05", "8", "", "수료", "", ""],
    ])
    res = svc.upload_records(db_session, content)
    assert res.created == 1
    rows = svc.list_records(db_session, employee_id=emp.id)
    assert len(rows) == 1
    assert rows[0].course_code == "C-LAW"


def test_upload_resolves_course_code(db_session, emp, course_law):
    content = _build_xlsx([
        _HEADER,
        ["EDU001", "이수자", "C-LAW", "법정필수교육",
         "20250601", "20250605", "8.0", "85", "수료", "CERT-A", "OK"],
    ])
    res = svc.upload_records(db_session, content)
    assert res.created == 1
    row = svc.list_records(db_session)[0]
    assert row.course_id == course_law.id
    assert row.start_date == date(2025, 6, 1)
    assert row.hours == Decimal("8.0")
    assert row.score == Decimal("85")
    assert row.certificate_no == "CERT-A"


def test_upload_unknown_course_code_creates_without_course(db_session, emp):
    content = _build_xlsx([
        _HEADER,
        ["EDU001", "이수자", "UNKNOWN-CODE", "신규과정",
         "2025-06-01", "", "4", "", "수료", "", ""],
    ])
    res = svc.upload_records(db_session, content)
    assert res.created == 1
    row = svc.list_records(db_session)[0]
    assert row.course_id is None
    assert row.course_name_snapshot == "신규과정"
