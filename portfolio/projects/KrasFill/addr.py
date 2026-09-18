# -*- coding: utf-8 -*-
"""파일명/문자열에서 주소(시도·시군구·읍면동·지번 목록)를 파싱한다.

HUG 감정평가 워크북 파일명 관례:
  '(복사본)26_1-경기 부천시 원미구 원미동148-21, 148-22, ... 청담빌리지 3층 301호 (...)_...'
"""
import re

SIDO_NAMES = [
    "서울", "부산", "대구", "인천", "광주", "대전", "울산", "세종",
    "경기", "강원", "충북", "충남", "전북", "전남", "경북", "경남", "제주",
]

_ADDR_RE = re.compile(
    r"(서울|부산|대구|인천|광주|대전|울산|세종|경기|강원|충북|충남|전북|전남|경북|경남|제주)"
    r"(?:특별시|광역시|특별자치시|특별자치도|남도|북도|도|시)?\s*"          # 시도 접미
    r"((?:[가-힣]+(?:시|군|구)\s*)+)"                                      # 시군구 (복합 가능)
    r"([가-힣0-9]+?(?:동|가|읍|면|리))\s*"                                  # 읍면동
    # 지번 목록. 층/호 앞 숫자('5층504호'의 5)를 지번으로 오인하지 않도록 제외
    r"(산?\d+(?:-\d+)?(?![층호])(?:\s*,\s*산?\d+(?:-\d+)?(?![층호]))*)"

)


def parse_address(text):
    """주소 문자열/파일명에서 (시도, 시군구, 읍면동, [지번,...]) 추출.

    Returns dict or None.
    """
    m = _ADDR_RE.search(text)
    if not m:
        return None
    sido, sgg, umd, jibun_part = m.groups()
    jibuns = [j.strip() for j in jibun_part.split(",") if j.strip()]
    return {
        "sido": sido,
        "sgg": re.sub(r"\s+", "", sgg),   # '부천시 원미구' -> '부천시원미구'
        "umd": umd,
        "jibuns": jibuns,
    }


def parse_address_from_workbook(xlsx_path):
    """파일명에 주소가 없을 때: 워크북 내부 헤더 셀에서 주소를 찾는다.

    (건축대장) K4, (토지이용) I3 에 전체 제목(주소 포함)이 들어 있다.
    """
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    candidates = [("(건축대장)", "K4"), ("(토지이용)", "I3")]
    try:
        for sheet, cell in candidates:
            if sheet not in wb.sheetnames:
                continue
            v = wb[sheet][cell].value
            if v:
                info = parse_address(str(v))
                if info:
                    return info
        # 최후: 주요 시트 상단 몇 행을 훑는다
        for sheet in ("(건축대장)", "(토지이용)", "(정식)", "의뢰및유의사항"):
            if sheet not in wb.sheetnames:
                continue
            ws = wb[sheet]
            for row in ws.iter_rows(min_row=1, max_row=10):
                for c in row:
                    if isinstance(c.value, str) and len(c.value) > 10:
                        info = parse_address(c.value)
                        if info:
                            return info
    finally:
        wb.close()
    return None


def read_parcels_from_jsik(xlsx_path):
    """(정식) 탭 토지 목록(F열)에서 지번 리스트를 읽는다. 없으면 []"""
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    jibuns = []
    try:
        if "(정식)" not in wb.sheetnames:
            return []
        ws = wb["(정식)"]
        for row in ws.iter_rows(min_row=3, max_row=12, min_col=6, max_col=6):
            v = row[0].value
            if v and re.match(r"^산?\d+(-\d+)?$", str(v).strip()):
                jibuns.append(str(v).strip())
    finally:
        wb.close()
    return jibuns


def read_parcels_from_deunggi(xlsx_path):
    """(등기) 탭 하단 토지 목록(F열)에서 전체 지번을 읽는다. 없으면 []

    등기부 표제부의 '대지권 목적인 토지' 목록이라 '외 N필지' 건의 전체 필지 원천.
    앵커: B열 '토지' + G열 '지목' 헤더 행을 찾아 그 아래 F열을 읽는다.
    """
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        if "(등기)" not in wb.sheetnames:
            return []
        ws = wb["(등기)"]
        rows = list(ws.iter_rows(min_row=1, max_row=100, min_col=2, max_col=7,
                                 values_only=True))
        anchor = None
        for i, row in enumerate(rows):
            b = str(row[0]).strip() if row[0] is not None else ""
            g = str(row[5]).strip() if row[5] is not None else ""
            if b == "토지" and g == "지목":
                anchor = i
                break
        if anchor is None:
            return []
        jibuns = []
        for row in rows[anchor + 1:anchor + 13]:  # 토지이용 블록 최대 10개 + 여유
            f = str(row[4]).strip() if row[4] is not None else ""
            if re.match(r"^산?\d+(-\d+)?$", f):
                jibuns.append(f)
            elif jibuns:  # 목록이 시작된 뒤 빈 행/합계 행이 나오면 끝
                break
        return jibuns
    finally:
        wb.close()


def jibun_to_bobn_bubn(jibun):
    """'148-21' -> (mountain, '0148', '0021'). '산12' -> (True, '0012', '0000')"""
    jibun = jibun.strip()
    mountain = jibun.startswith("산")
    if mountain:
        jibun = jibun[1:].strip()
    parts = jibun.split("-")
    bobn = parts[0].strip().zfill(4)
    bubn = (parts[1].strip() if len(parts) > 1 else "0").zfill(4)
    return mountain, bobn, bubn
