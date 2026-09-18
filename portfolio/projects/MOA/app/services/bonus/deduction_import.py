"""공제 대장 엑셀 올리기 — 재무팀이 표로 적어 온 공제를 한 번에 등록한다.

양식(첫 줄 머리글): 사람 | 종류 | 금액 | 감정서 | 발생일 | 메모 | 적용월
- 종류는 한글(화환·패널티·보험료·기타비용·감정서경비·감정서경비환입·선지급·기타공제·처리비·가변비추가·가변비환입·미납비이월·기타)
  또는 코드(WREATH …) 둘 다 받는다.
- 적용월(YYYYMM 또는 YYYY-MM)을 적으면 그 달(열려 있어야 함)에 바로 적용, 비우면 대기.
- 같은 파일을 두 번 올려도 두 번 들어가지 않는다 — 행마다 파일 해시+행번호를 출처 키로 쓴다.
"""

import hashlib
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook

from app.services.bonus.engine import DEDUCTION_KINDS

KIND_LABELS = {
    "화환": "WREATH", "화환공제": "WREATH", "패널티": "PENALTY", "벌과금": "PENALTY", "보험료": "INSURANCE", "법인차량": "INSURANCE",
    "자동차세": "INSURANCE", "기타비용": "OTHER_EXPENSE", "기타 비용": "OTHER_EXPENSE", "감정서경비": "DOC_EXPENSE",
    "당월감정서경비": "DOC_EXPENSE", "감정서경비환입": "EXPENSE_CREDIT", "선지급": "ADVANCE_PAID", "기타공제": "OTHER_DEDUCT",
    "미수금": "OTHER_DEDUCT", "미수금회수": "OTHER_DEDUCT",

    "처리비": "HANDLING", "처리수당": "HANDLING", "가변비추가": "VARIABLE_ADJ", "가변비": "VARIABLE_ADJ", "가변비환입": "VARIABLE_CREDIT",
    "미납비이월": "UNPAID_CARRY", "기타": "MISC",
}
HEADERS = {
    "person": ("사람", "이름", "성명", "담당자"), "kind": ("종류", "구분", "공제종류"), "amount": ("금액",),
    "doc_id": ("감정서", "감정서번호"), "occurred_on": ("발생일", "날짜", "일자"), "memo": ("메모", "내역", "비고"),
    "apply_period": ("적용월", "지급월"),
}
TEMPLATE_ROWS = [
    ("강무진", "화환", 80000, "", "2026-08-12", "○○ 부친상", ""),
    ("김정원", "감정서경비", 20000, "01-2606-5-0085", "2026-08-15", "전자수입인지", "202609"),
    ("고세욱", "선지급", 4610400, "01-2505-4-0176", "", "제이앤파트너스 선지급 상여", ""),
]


def _kind(value: Any) -> str:
    text = str(value or "").strip().replace(" ", "")
    if text.upper() in DEDUCTION_KINDS:
        return text.upper()
    return KIND_LABELS.get(text, "")


def _date(value: Any) -> "date | None":
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace(".", "-").replace("/", "-")
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _period(value: Any) -> "str | None":
    text = str(value or "").strip().replace("-", "").replace(".", "")
    if isinstance(value, (date, datetime)):
        text = f"{value:%Y%m}"
    if not text:
        return None
    if len(text) == 6 and text.isdigit() and 1 <= int(text[4:]) <= 12:
        return text
    return "INVALID"


def parse_workbook(content: bytes, *, source_label: str = "xlsx") -> "tuple[list[dict[str, Any]], list[str]]":
    """첫 시트 → (등록할 행 목록, 오류 목록). 오류 행은 건너뛰고 나머지는 그대로 돌려준다."""
    digest = hashlib.sha1(content).hexdigest()[:10]
    wb = load_workbook(BytesIO(content), data_only=True)
    ws = wb.worksheets[0]
    header = [str(ws.cell(1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
    columns: "dict[str, int]" = {}
    for key, names in HEADERS.items():
        for index, name in enumerate(header, start=1):
            if name in names:
                columns[key] = index
                break
    missing = [key for key in ("person", "kind", "amount") if key not in columns]
    if missing:
        return [], [f"머리글에 {', '.join('사람' if m == 'person' else '종류' if m == 'kind' else '금액' for m in missing)} 열이 없습니다 (첫 줄: {' | '.join(h for h in header if h)})"]
    rows: "list[dict[str, Any]]" = []
    errors: "list[str]" = []
    for r in range(2, ws.max_row + 1):
        cell = lambda key: ws.cell(r, columns[key]).value if key in columns else None  # noqa: E731
        person = str(cell("person") or "").strip()
        if not person and cell("amount") in (None, "") and not str(cell("kind") or "").strip():
            continue                                                     # 빈 줄
        kind = _kind(cell("kind"))
        try:
            amount = float(str(cell("amount") or 0).replace(",", ""))
        except ValueError:
            amount = -1
        period = _period(cell("apply_period"))
        problems = []
        if not person:
            problems.append("사람 없음")
        if not kind:
            problems.append(f"종류를 모름({cell('kind')!r})")
        if amount <= 0:
            problems.append("금액이 0 이하")
        if period == "INVALID":
            problems.append(f"적용월 형식 오류({cell('apply_period')!r})")
        if problems:
            errors.append(f"{r}행: {', '.join(problems)}")
            continue
        rows.append({
            "person": person, "kind": kind, "amount": amount,
            "doc_id": (str(cell("doc_id") or "").strip() or None),
            "occurred_on": _date(cell("occurred_on")), "memo": (str(cell("memo") or "").strip() or None),
            "source": "MANUAL", "source_key": f"{source_label}:{digest}:{r}", "apply_period": period,
        })
    return rows, errors


def template_workbook() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "공제"
    ws.append(["사람", "종류", "금액", "감정서", "발생일", "메모", "적용월"])
    for row in TEMPLATE_ROWS:
        ws.append(list(row))
    for column, width in zip("ABCDEFG", (12, 16, 14, 18, 14, 36, 10)):
        ws.column_dimensions[column].width = width
    guide = wb.create_sheet("종류")
    guide.append(["종류(한글)", "코드", "뜻"])
    for label, code, meaning in (
        ("화환", "WREATH", "화환공제(W) — 상여 0이면 다음 달 이월"), ("패널티", "PENALTY", "은행 패널티(W)"),
        ("보험료", "INSURANCE", "보험료·법인차량·자동차세(W)"), ("기타비용", "OTHER_EXPENSE", "그 밖의 비용(W)"),
        ("감정서경비", "DOC_EXPENSE", "당월감정서경비(O)"), ("감정서경비환입", "EXPENSE_CREDIT", "앞 달 감정서경비 돌려줌(+)"),
        ("선지급", "ADVANCE_PAID", "세전상여에서 뺌"), ("미수금", "OTHER_DEDUCT", "미수금 회수 — 세금 다 뗀 지급액에서 뺌(AB·세후)"),
        ("처리비", "HANDLING", "처리비·처리수당(+)"), ("가변비추가", "VARIABLE_ADJ", "가변비 더함"),
        ("가변비환입", "VARIABLE_CREDIT", "가변비 돌려줌(+)"), ("미납비이월", "UNPAID_CARRY", "미납비이월 수기"), ("기타", "MISC", "기타(N)"),
    ):
        guide.append([label, code, meaning])
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["C"].width = 40
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
