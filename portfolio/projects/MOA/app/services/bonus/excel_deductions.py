"""엑셀 공제내역 두 시트 → 공제 대장 (a10_bonus_deduction, source='EXCEL').

재무팀 엑셀은 공제를 두 시트에 쌓는다.
- '감정서관련비용공제내역' : 전자수입인지·제본·협회비 … (A 날짜, C 내역, D 담당자, E 금액, G 종류) → DOC_EXPENSE(O)
- '화환및감정서관련없는비용공제내역' : 화환·패널티·법인차량 … (A 날짜, B 담당자, C 내역, D 금액, E 종류) → 화환류(W)
각 달의 몫은 그 달 주주 시트 합계행의 O/W 수식이 가리키는 **행 범위**(월 블록)다 — 라벨이 아니라 수식이 진실.
W 는 IF(Q<0, 0, …) 이라 상여가 0인 사람의 화환류는 그 달에 빠지지 않는다 → 대기(PENDING)로 남긴다.
블록 밖(맨 뒤 '공제못함' 등)의 행도 대기. 선지급은 합계행 Y 수식 끝의 상수, 기타공제(AB)는 합계행 값.
"""

import re
from datetime import date, datetime
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete

from app.models.bonus_ledger import BonusDeduction
from app.services.bonus.rules import extract_doc_ids

DOC_SHEET = "감정서관련비용공제내역"
PERSON_SHEET = "화환및감정서관련없는비용공제내역"
_MONTH = re.compile(r"^(\d\d)\.(\d\d)월-주주$")
_RANGE = re.compile(r"\$[A-Z]+\$(\d+):\$[A-Z]+\$(\d+)")
_ADVANCE = re.compile(r"-\s*(\d[\d,]*)\s*$")
_KIND_RULES = (
    ("화환", "WREATH"), ("패널티", "PENALTY"), ("보험", "INSURANCE"), ("법인차량", "INSURANCE"), ("자동차세", "INSURANCE"),
)


def map_kind(kind_text: str) -> str:
    text = str(kind_text or "")
    for needle, kind in _KIND_RULES:
        if needle in text:
            return kind
    return "OTHER_EXPENSE"


def _number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _date(value: Any) -> "date | None":
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if match:
        return date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    return None


def _period(sheet: str) -> str:
    match = _MONTH.match(sheet)
    return f"20{match.group(1)}{match.group(2)}"


def _total_rows(ws) -> "list[int]":
    return [r for r in range(1, ws.max_row + 1) if "합계" in str(ws.cell(r, 1).value or "")]


def read_month_blocks(wb_f) -> "dict[str, dict[str, tuple[int, int] | None]]":
    """{지급월: {"doc": (start, end), "person": (start, end)}} — 합계행 O(15)/W(23) 수식의 행 범위."""
    blocks: "dict[str, dict[str, tuple[int, int] | None]]" = {}
    for name in wb_f.sheetnames:
        if not _MONTH.match(name):
            continue
        ws = wb_f[name]
        found: "dict[str, tuple[int, int] | None]" = {"doc": None, "person": None}
        for r in _total_rows(ws):
            for key, col, sheet in (("doc", 15, DOC_SHEET), ("person", 23, PERSON_SHEET)):
                value = ws.cell(r, col).value
                if found[key] is None and isinstance(value, str) and sheet in value:
                    match = _RANGE.search(value)
                    if match:
                        found[key] = (int(match.group(1)), int(match.group(2)))
            if found["doc"] and found["person"]:
                break
        if found["doc"] or found["person"]:
            blocks[_period(name)] = found
    return blocks


def _new_total() -> "dict[str, Any]":
    return {"wreath": 0.0, "advances": [], "other_deduct": 0.0, "other_formula": None, "in_shareholder": False, "in_associate": False}


def read_person_totals(wb_f, wb_v, sheet: str) -> "dict[str, dict[str, Any]]":
    """합계행에서 사람별 W(화환 적용액)·선지급(Y 수식 상수)·AB(기타공제)를 읽는다. 블록이 여럿이면 합친다.

    같은 달 평·동 시트(소속·공통건)의 AC(화환)도 더한다 — 소속평가사의 화환은 거기서 빠진다.
    """
    ws_f, ws_v = wb_f[sheet], wb_v[sheet]
    totals: "dict[str, dict[str, Any]]" = {}
    associate_sheet = sheet.replace("-주주", "-평.동")
    if associate_sheet in wb_v.sheetnames:
        pd = wb_v[associate_sheet]
        for r in range(1, pd.max_row + 1):
            owner = str(pd.cell(r, 4).value or "").strip()
            if "합계" in str(pd.cell(r, 3).value or "") and owner:
                entry = totals.setdefault(owner, _new_total())
                entry["in_associate"] = True
                entry["wreath"] += _number(pd.cell(r, 29).value)
    for r in _total_rows(ws_v):
        owner = str(ws_v.cell(r, 4).value or "").strip()
        if not owner:
            continue
        entry = totals.setdefault(owner, _new_total())
        entry["in_shareholder"] = True
        entry["wreath"] += _number(ws_v.cell(r, 23).value)
        y_formula = ws_f.cell(r, 25).value
        if isinstance(y_formula, str):
            match = _ADVANCE.search(y_formula.replace(" ", ""))
            if match and "ROUNDDOWN" in y_formula:
                entry["advances"].append((r, float(match.group(1).replace(",", ""))))
        other = _number(ws_v.cell(r, 28).value)
        if other:
            entry["other_deduct"] += other
            formula = ws_f.cell(r, 28).value
            entry["other_formula"] = str(formula) if formula is not None else None
    return totals


def _doc_rows(ws, start: int, end: int, sheet_label: str) -> "list[dict[str, Any]]":
    rows = []
    for r in range(start, end + 1):
        person = str(ws.cell(r, 4).value or "").strip()
        amount = _number(ws.cell(r, 5).value)
        if not person or amount == 0:
            continue                                                     # 소계·빈 줄
        detail = str(ws.cell(r, 3).value or "").strip()
        kind_text = str(ws.cell(r, 7).value or "").strip()
        docs = extract_doc_ids(detail)
        rows.append({
            "person": person, "kind": "DOC_EXPENSE" if amount > 0 else "EXPENSE_CREDIT",   # 음수 = 앞 달에 뺀 것을 돌려줌
            "amount": abs(amount), "doc_id": docs[0] if docs else None,
            "occurred_on": _date(ws.cell(r, 1).value), "memo": " ".join(x for x in (kind_text, detail) if x) or None,
            "source": "EXCEL", "source_key": f"doc:{sheet_label}:{r}", "row": r,
        })
    return rows


def _person_rows(ws, start: int, end: int, sheet_label: str) -> "list[dict[str, Any]]":
    rows = []
    for r in range(start, end + 1):
        person = str(ws.cell(r, 2).value or "").strip()
        amount = _number(ws.cell(r, 4).value)
        if not person or amount <= 0:
            continue
        detail = str(ws.cell(r, 3).value or "").strip()
        kind_text = str(ws.cell(r, 5).value or "").strip()
        docs = extract_doc_ids(detail)
        rows.append({
            "person": person, "kind": map_kind(kind_text), "amount": amount, "doc_id": docs[0] if docs else None,
            "occurred_on": _date(ws.cell(r, 1).value), "memo": " ".join(x for x in (kind_text, detail) if x) or None,
            "kind_text": kind_text, "source": "EXCEL", "source_key": f"wr:{sheet_label}:{r}", "row": r,
        })
    return rows


def _month_label(period: str) -> str:
    return f"{period[2:4]}.{period[4:6]}"


def build_items(wb_f, wb_v, *, periods: "list[str] | None" = None) -> "list[dict[str, Any]]":
    """두 시트 + 합계행 → 대장 항목 목록 (status/applied_period 포함). 블록 밖 행은 대기."""
    blocks = read_month_blocks(wb_f)
    items: "list[dict[str, Any]]" = []
    doc_ws, person_ws = wb_v[DOC_SHEET], wb_v[PERSON_SHEET]
    covered_doc: "set[int]" = set()
    covered_person: "set[int]" = set()
    applied_person_keys: "set[tuple]" = set()
    for period in sorted(blocks):
        label = _month_label(period)
        wanted = periods is None or period in periods
        doc_range, person_range = blocks[period]["doc"], blocks[period]["person"]
        totals = read_person_totals(wb_f, wb_v, f"{label}월-주주")
        if doc_range:
            covered_doc.update(range(doc_range[0], doc_range[1] + 1))
            if wanted:
                for row in _doc_rows(doc_ws, *doc_range, sheet_label=label):
                    if totals.get(row["person"], {}).get("in_shareholder"):      # O 는 주주 시트 합계행에서만 빠진다
                        items.append({**row, "status": "APPLIED", "applied_period": period})
                    else:
                        items.append({**row, "status": "PENDING", "applied_period": None,
                                      "memo": f"{row['memo'] or ''} (엑셀 {label} 주주 시트에 없음)".strip()})
        if person_range:
            covered_person.update(range(person_range[0], person_range[1] + 1))
            for row in _person_rows(person_ws, *person_range, sheet_label=label):
                applied = totals.get(row["person"], {}).get("wreath", 0.0) > 0
                if not wanted:
                    continue
                if applied:
                    applied_person_keys.add((row["person"], row["occurred_on"], row["amount"], row["kind_text"]))
                    items.append({**row, "status": "APPLIED", "applied_period": period})
                else:
                    items.append({**row, "status": "PENDING", "applied_period": None,
                                  "memo": f"{row['memo'] or ''} (엑셀 {label} Q<0 공제 못함)".strip()})
        if wanted:
            for owner, total in totals.items():
                for row, amount in total["advances"]:
                    items.append({
                        "person": owner, "kind": "ADVANCE_PAID", "amount": amount, "doc_id": None, "occurred_on": None,
                        "memo": "엑셀 Y 선지급", "source": "EXCEL", "source_key": f"adv:{label}:{row}",
                        "status": "APPLIED", "applied_period": period,
                    })
                if total["other_deduct"]:
                    items.append({
                        "person": owner, "kind": "OTHER_DEDUCT", "amount": total["other_deduct"], "doc_id": None, "occurred_on": None,
                        "memo": f"엑셀 AB {total['other_formula'] or ''}".strip(), "source": "EXCEL", "source_key": f"ab:{label}:{owner}",
                        "status": "APPLIED", "applied_period": period,
                    })
    # 어느 블록에도 안 든 행(맨 뒤 '공제못함', 다음 달 준비분)은 대기. 앞 블록에 이미 적용된 것의 복사본이면 건너뛴다.
    last_doc = max(covered_doc) if covered_doc else 0
    last_person = max(covered_person) if covered_person else 0
    for row in _doc_rows(doc_ws, last_doc + 1, doc_ws.max_row, sheet_label="tail"):
        items.append({**row, "status": "PENDING", "applied_period": None})
    for row in _person_rows(person_ws, last_person + 1, person_ws.max_row, sheet_label="tail"):
        if (row["person"], row["occurred_on"], row["amount"], row["kind_text"]) in applied_person_keys:
            continue
        items.append({**row, "status": "PENDING", "applied_period": None})
    return [{k: v for k, v in item.items() if k not in ("row", "kind_text")} for item in items]


def seed_deductions(
    db, workbook_path, *, periods: "list[str] | None" = None, dry_run: bool = False, force: bool = False, workbooks=None,
) -> "dict[str, Any]":
    """엑셀 → 대장. 같은 source_key 는 건너뛴다(force 면 EXCEL 출처를 지우고 다시 넣는다). 마감된 달에도 이력으로 넣는다."""
    from app.services.bonus.deduction_items import insert_seeded

    if workbooks is None:
        workbooks = (load_workbook(workbook_path, data_only=False), load_workbook(workbook_path, data_only=True))
    items = build_items(*workbooks, periods=periods)
    report: "dict[str, Any]" = {"dry_run": dry_run, "written": 0, "skipped": 0, "by_period": {}, "pending_people": {}}
    for item in items:
        key = item["applied_period"] or "PENDING"
        report["by_period"][key] = report["by_period"].get(key, 0) + 1
        if item["status"] == "PENDING":
            report["pending_people"][item["person"]] = report["pending_people"].get(item["person"], 0) + item["amount"]
    if dry_run:
        report["written"] = len(items)
        return report
    if force:
        db.execute(delete(BonusDeduction).where(BonusDeduction.source == "EXCEL"))
        db.commit()
    written, skipped = insert_seeded(db, items)
    report["written"], report["skipped"] = written, skipped
    return report
