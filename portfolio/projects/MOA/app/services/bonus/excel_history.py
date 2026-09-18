"""엑셀로 지급한 달을 EXCEL 마감 이력으로 넣는다 (2단계 16).

왜: 새 엔진은 마감된 결과(a10_bonus_result)로 (1) 같은 감정서를 두 달에 두 번 주지 않고
(2) 전월 미납비용을 이번 달 미납비이월로 가져온다. 도입 전 달은 엑셀에만 있으니 그 시트를
그대로 스냅샷으로 옮긴다. 행의 fee 는 사람 몫 산정금액(H)이다 — 기지급 차감의 기준.
MOA 로 마감한 달(source='MOA')은 건드리지 않는다.
"""

import re
from datetime import datetime
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import delete, select

from app.models.bonus_ledger import BonusClose, BonusResult
from app.services.bonus.excel_compare import read_associate_sheet, read_shareholder_sheet

_SHEET = re.compile(r"^(\d\d)\.(\d\d)월-주주$")


def _period(sheet: str) -> str:
    match = _SHEET.match(sheet)
    return f"20{match.group(1)}{match.group(2)}"


def _rows_for(period: str, shareholder: "dict[str, Any]", associate: "dict[str, Any]") -> "list[BonusResult]":
    rows: "list[BonusResult]" = []
    for (doc, owner), entry in shareholder["rows"].items():
        rows.append(BonusResult(
            period=period, person=owner, kind="SHAREHOLDER", doc_id=doc,
            fee=entry["assessed"], assessed=entry["assessed"], source="EXCEL",
        ))
    for owner, totals in shareholder["totals"].items():
        rows.append(BonusResult(
            period=period, person=owner, kind="SHAREHOLDER", doc_id=None,
            payout_base=totals.get("payout_base", 0.0), bonus=totals.get("bonus", 0.0),
            pretax=totals.get("pretax", 0.0), income_tax=totals.get("income_tax", 0.0),
            resident_tax=totals.get("resident_tax", 0.0), other_deduct=totals.get("other_deduct", 0.0),
            payment=totals.get("payment", 0.0), unpaid_carry_out=totals.get("unpaid_carry_out", 0.0),
            card_limit=totals.get("card_limit", 0.0), source="EXCEL",
        ))
    kinds = {person: totals["kind"] for person, totals in associate["totals"].items()}
    for (doc, person), entry in associate["rows"].items():
        rows.append(BonusResult(
            period=period, person=person, kind=kinds.get(person, "ASSOCIATE"), doc_id=doc,
            fee=entry["assessed"], assessed=entry["assessed"], bonus=entry["bonus"], source="EXCEL",
        ))
    for person, totals in associate["totals"].items():
        rows.append(BonusResult(
            period=period, person=person, kind=totals["kind"], doc_id=None,
            bonus=totals["bonus"], pretax=totals["pretax"], income_tax=totals["income_tax"],
            resident_tax=totals["resident_tax"], other_deduct=totals["other_deduct"],
            payment=totals["payment"], unpaid_carry_out=0.0, source="EXCEL",
        ))
    return rows


def seed_history(
    db, workbook_path, *, periods: "list[str] | None" = None, dry_run: bool = False, force: bool = False,
) -> "dict[str, Any]":
    """워크북의 'YY.MM월-주주'(+평.동) 시트를 EXCEL 마감으로. 무엇을 했는지 돌려준다."""
    wb = load_workbook(workbook_path, data_only=True)
    report: "dict[str, Any]" = {"dry_run": dry_run, "written": [], "skipped": [], "protected": [], "rows": {}}
    for sheet in sorted(name for name in wb.sheetnames if _SHEET.match(name)):
        period = _period(sheet)
        if periods and period not in periods:
            continue
        existing = db.get(BonusClose, period)
        if existing is not None and existing.source != "EXCEL":
            report["protected"].append(period)
            continue
        if existing is not None and not force:
            report["skipped"].append(period)
            continue
        shareholder = read_shareholder_sheet(wb[sheet])
        associate_sheet = sheet.replace("-주주", "-평.동")
        associate = read_associate_sheet(wb[associate_sheet]) if associate_sheet in wb.sheetnames else {"rows": {}, "totals": {}}
        rows = _rows_for(period, shareholder, associate)
        report["written"].append(period)
        report["rows"][period] = len(rows)
        if dry_run:
            continue
        db.execute(delete(BonusResult).where(BonusResult.period == period, BonusResult.source == "EXCEL"))
        if existing is None:
            db.add(BonusClose(period=period, status="CLOSED", source="EXCEL", closed_at=datetime.now(), memo=f"엑셀 {sheet} 이력"))
        else:
            existing.status, existing.closed_at, existing.memo = "CLOSED", datetime.now(), f"엑셀 {sheet} 이력(다시 씀)"
        db.add_all(rows)
        db.commit()
    return report


def closed_periods(db) -> "list[str]":
    return sorted(row.period for row in db.scalars(select(BonusClose).where(BonusClose.status == "CLOSED")).all())
