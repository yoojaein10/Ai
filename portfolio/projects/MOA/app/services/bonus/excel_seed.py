"""재무팀 성과상여 엑셀 → 마스터 초기값 읽기 (순수 함수, DB 없음).

시트 이름 'YY.MM월-주주' / 'YY.MM월-평.동' 만 읽는다(정렬하면 시간순).
- 요율: 주주 시트 '합계' 행의 C 라벨 + R/S/T/U 중 `*$R$6` 꼴 수식이 있는 열(40/45/35/30).
  사람별로 **가장 최근 시트**의 블록을 쓴다.
- 지분: 상세행(A가 감정서번호)을 다음 합계행의 주인(D)에게 귀속한다. H 수식의 `*25%`
  가 지분이고, 수식이 없는 달(25.12·26.01)은 값 H/(F+G) 로 읽는다. 혼자 100% 인
  감정서는 저장하지 않는다.
- 소속 파라미터: 평.동 시트 '소속 합계' 행의 AD(`*70%` = 지급률)·AE(`*15%` = 소득세율).
"""

import re
from typing import Any

from openpyxl import load_workbook

from app.services.bonus.shares import match_note_names, parse_share_note

SHAREHOLDER_SHEET = re.compile(r"^\d\d\.\d\d월-주주$")
ASSOCIATE_SHEET = re.compile(r"^\d\d\.\d\d월-평\.동$")
RATE_COLUMNS = {18: 40, 19: 45, 20: 35, 21: 30}  # R S T U
_RATE_FORMULA = re.compile(r"\*\s*\$[RSTU]\$6")
_DOC = re.compile(r"^\d\d-\d{4}")
_PCT_FORMULA = re.compile(r"\*\s*([\d.]+)\s*%")
_HALF_FORMULA = re.compile(r"/\s*2\b|\*\s*0\.5\b")
DEFAULT_ASSOCIATE_TAX = 0.15  # 26.08 시트 소속 25명 중 22명


def open_formulas(path: str) -> Any:
    return load_workbook(path, data_only=False)


def open_values(path: str) -> Any:
    return load_workbook(path, data_only=True)


def _month_sheets(wb: Any, pattern: "re.Pattern[str]") -> "list[str]":
    return sorted(name for name in wb.sheetnames if pattern.match(name))


def _is_total(value: Any) -> bool:
    return isinstance(value, str) and "합계" in value


def _is_doc(value: Any) -> bool:
    return isinstance(value, str) and bool(_DOC.match(value))


def _sheet_key(sheet: str) -> str:
    return sheet[:5]  # 'YY.MM'


def read_rate_blocks(wb: Any) -> "dict[str, dict[str, Any]]":
    """{사람: {sheet, blocks: [{label, rate|None}]}} — 가장 최근 시트가 이긴다."""
    result: "dict[str, dict[str, Any]]" = {}
    for sheet in _month_sheets(wb, SHAREHOLDER_SHEET):
        ws = wb[sheet]
        per_person: "dict[str, list[dict[str, Any]]]" = {}
        for r in range(7, ws.max_row + 1):
            if not _is_total(ws.cell(r, 1).value):
                continue
            person = str(ws.cell(r, 4).value or "").strip()
            if not person:
                continue
            label = str(ws.cell(r, 3).value or "").strip()
            rate = next(
                (pct for col, pct in RATE_COLUMNS.items()
                 if isinstance(ws.cell(r, col).value, str) and _RATE_FORMULA.search(ws.cell(r, col).value)),
                None,
            )
            per_person.setdefault(person, []).append({"label": label, "rate": rate})
        for person, blocks in per_person.items():
            result[person] = {"sheet": sheet, "blocks": blocks}
    return result


def _share_pct(ws_f: Any, ws_v: Any, r: int) -> "float | None":
    formula = ws_f.cell(r, 8).value
    if isinstance(formula, str) and formula.startswith("="):
        m = _PCT_FORMULA.search(formula)
        if m:
            return float(m.group(1))
        return 50.0 if _HALF_FORMULA.search(formula) else 100.0
    assessed = ws_v.cell(r, 8).value
    base = float(ws_v.cell(r, 6).value or 0) + float(ws_v.cell(r, 7).value or 0)
    if assessed is None or not base:
        return None
    return round(float(assessed) / base * 100, 3)


def read_share_rows(wb_f: Any, wb_v: Any) -> "list[dict[str, Any]]":
    """(감정서, 블록 주인) 별 지분 — 혼자 100% 인 감정서는 뺀다."""
    latest: "dict[tuple[str, str], dict[str, Any]]" = {}
    owners: "dict[str, set[str]]" = {}
    for sheet in _month_sheets(wb_f, SHAREHOLDER_SHEET):
        ws_f, ws_v = wb_f[sheet], wb_v[sheet]
        pending: "list[int]" = []
        for r in range(7, ws_f.max_row + 1):
            head = ws_f.cell(r, 1).value
            if _is_doc(head):
                pending.append(r)
                continue
            if not _is_total(head):
                continue
            owner = str(ws_f.cell(r, 4).value or "").strip()
            for row in pending:
                pct = _share_pct(ws_f, ws_v, row)
                if not owner or pct is None or pct <= 0:
                    continue
                doc = str(ws_f.cell(row, 1).value).split(",")[0].strip()
                customer = str(ws_f.cell(row, 3).value or "").strip()
                note = parse_share_note(customer)
                bc = match_note_names(note.bc, [owner]).get(owner) if note and note.bc else None
                latest[(doc, owner)] = {
                    "doc_id": doc, "person": owner, "share_pct": pct, "bc_pct": bc,
                    "note": customer[:200] or None, "sheet": sheet,
                }
                owners.setdefault(doc, set()).add(owner)
            pending = []
    return [
        row for (doc, _), row in latest.items()
        if len(owners[doc]) > 1 or row["share_pct"] != 100 or row["bc_pct"] is not None
    ]


def read_associate_params(wb: Any) -> "dict[str, dict[str, Any]]":
    """'소속 합계' 행의 AD(지급률)·AE(소득세율) 수식 — 가장 최근 시트가 이긴다."""
    result: "dict[str, dict[str, Any]]" = {}
    for sheet in _month_sheets(wb, ASSOCIATE_SHEET):
        ws = wb[sheet]
        for r in range(8, ws.max_row + 1):
            if str(ws.cell(r, 3).value or "").strip() != "소속 합계":
                continue
            person = str(ws.cell(r, 4).value or "").strip()
            if not person:
                continue
            pay = _PCT_FORMULA.search(str(ws.cell(r, 30).value or ""))
            tax = _PCT_FORMULA.search(str(ws.cell(r, 31).value or ""))
            result[person] = {
                "pay_ratio": float(pay.group(1)) / 100 if pay else 1.0,
                "tax_rate": float(tax.group(1)) / 100 if tax else DEFAULT_ASSOCIATE_TAX,
                "sheet": sheet,
            }
    return result


def read_person_list(
    rate_blocks: "dict[str, dict[str, Any]]", associate_params: "dict[str, dict[str, Any]]"
) -> "list[dict[str, Any]]":
    """주주(요율 블록이 있는 사람) + 소속.

    양쪽에 다 있으면 소속 합계가 주주 블록과 같은 달이거나 더 최근이면 소속이다 — 옛 주주
    블록이 남은 이영은('~2023.12'), 이제 막 주주가 된 정인수('2026년 7월~' + 소속 합계)가
    그렇다. 접수일이 요율 구간에 드는 건은 리포트가 주주 행으로 올린다(schedule 우선).
    """
    people: "dict[str, dict[str, Any]]" = {
        person: {"person": person, "kind": "SHAREHOLDER", "pay_ratio": 1.0, "tax_rate": 0.30, "_sheet": info["sheet"]}
        for person, info in rate_blocks.items()
    }
    for person, params in associate_params.items():
        current = people.get(person)
        if current is None or _sheet_key(params["sheet"]) >= _sheet_key(current["_sheet"]):
            people[person] = {
                "person": person, "kind": "ASSOCIATE",
                "pay_ratio": params["pay_ratio"], "tax_rate": params["tax_rate"], "_sheet": params["sheet"],
            }
    return [
        {key: value for key, value in entry.items() if not key.startswith("_")}
        for entry in sorted(people.values(), key=lambda item: item["person"])
    ]
