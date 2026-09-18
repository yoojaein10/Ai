"""상여 리포트 → 엑셀 3시트 (주주 / 평·동 / 총괄표). 열 순서는 재무팀 시트 순서를 따른다."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.services.bonus.summary import journal_lines

SHAREHOLDER_COLUMNS = [
    ("person", "성명"), ("doc_id", "감정서번호"), ("work_type", "업무구분"), ("customer_name", "거래처"),
    ("receipt_date", "접수일"), ("rate", "요율(%)"), ("fee", "순수수료(F)"), ("land_net", "토지조사비-부족출장비(G)"),
    ("share_pct", "지분(%)"), ("assessed", "산정수수료(H)"), ("variable_cost", "가변비(J)"), ("indemnity", "손해배상충당금(K)"),
    ("carry_in", "미납비이월(M)"), ("misc", "기타(N)"), ("doc_expense", "당월감정서경비(O)"), ("association_fee", "협회비·공제료(P)"),
    ("payout_base", "산출액(Q)"), ("bonus", "상여(R)"), ("handling", "처리비(S)"), ("wreath", "화환공제(W)"),
    ("survey_fee", "물건조사비(X)"), ("pretax", "세전상여(Y)"), ("advance_paid", "선지급"), ("income_tax", "소득세(Z)"),
    ("resident_tax", "주민세(AA)"), ("other_deduct", "기타공제(AB)"), ("payment", "지급액(AC)"), ("unpaid_carry_out", "미납비용(AE)"),
    ("card_limit", "B.C"), ("flags", "표시"),
]
ASSOCIATE_COLUMNS = [
    ("person", "성명"), ("kind", "구분"), ("doc_id", "감정서번호"), ("work_type", "업무구분"), ("customer_name", "거래처"),
    ("receipt_date", "접수일"), ("fee", "순수수료(F)"), ("land_fee", "토지조사비(G)"), ("gross", "산정수수료(I)"),
    ("indemnity", "손해배상충당금(J)"), ("assessed", "산정액(L)"), ("applied_rate", "요율(%)"), ("bonus", "상여"),
    ("handling", "처리수당(AB)"), ("wreath", "화환공제(AC)"), ("pay_ratio", "지급률"), ("pretax", "세전상여(AD)"),
    ("income_tax", "소득세(AE)"), ("resident_tax", "지방세(AF)"), ("other_deduct", "기타공제(AG)"), ("payment", "지급액(AH)"),
    ("card_limit", "B.C"), ("flags", "표시"),
]
SUMMARY_COLUMNS = [
    ("name", "성명"), ("kinds", "구분"), ("pretax", "세전상여"), ("income_tax", "소득세"), ("resident_tax", "주민세"),
    ("other_deduct", "기타공제"), ("payment", "지급액"), ("card_limit", "B.C"), ("retired", "퇴사"),
]
_KIND_LABEL = {"SHAREHOLDER": "주주", "COMMON": "공통건", "ASSOCIATE": "소속"}


def _cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (list, tuple)):
        return ", ".join(str(v) for v in value)
    if isinstance(value, bool):
        return "Y" if value else ""
    return value


def _sheet(workbook: Workbook, title: str, columns, rows: "list[dict[str, Any]]", *, first: bool = False) -> None:
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title
    sheet.append([header for _, header in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([_cell(row.get(key)) for key, _ in columns])
        if row.get("_bold"):
            for cell in sheet[sheet.max_row]:
                cell.font = Font(bold=True)
    for index, (_, header) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(11, min(28, len(header) * 2 + 4))
    sheet.freeze_panes = "E2"


def _block_label(block: "dict[str, Any]") -> str:
    rate = f"{block['rate']:g}%" if block.get("rate") is not None else "요율 없음"
    start = block.get("block_from")
    if isinstance(start, (date, datetime)):
        start = f"{start:%Y-%m-%d}~"
    return f"합계 {rate}" + (f" ({start})" if start else "")


def shareholder_rows(entries: "list[dict[str, Any]]") -> "list[dict[str, Any]]":
    rows: "list[dict[str, Any]]" = []
    for entry in entries:
        totals = entry["totals"]
        for row in entry["rows"]:
            rows.append({**row, "land_net": float(row.get("land_fee") or 0) - float(row.get("travel_fee") or 0)})
        for block in entry.get("blocks", []):
            person_level = {
                "variable_cost": totals.get("variable_cost"), "carry_in": totals.get("carry_in"),
                "doc_expense": totals.get("doc_expense"), "misc": totals.get("misc"),
            } if block.get("main") else {}
            rows.append({
                "person": entry["name"], "doc_id": _block_label(block), "rate": block.get("rate"), **block, **person_level,
                "flags": ["퇴사"] if entry.get("retired") else [], "_bold": True,
            })
        if not entry.get("blocks"):
            rows.append({"person": entry["name"], "doc_id": "합계", **totals, "_bold": True})
    return rows


def associate_rows(entries: "list[dict[str, Any]]") -> "list[dict[str, Any]]":
    rows: "list[dict[str, Any]]" = []
    for entry in entries:
        for row in entry["rows"]:
            rows.append({**row, "kind": _KIND_LABEL.get(row.get("kind"), row.get("kind"))})
        rows.append({
            "person": entry["name"], "kind": _KIND_LABEL.get(entry["kind"], entry["kind"]), "doc_id": "합계",
            **entry["totals"], "flags": ["퇴사"] if entry.get("retired") else [], "_bold": True,
        })
    return rows


def summary_rows(summary: "list[dict[str, Any]]") -> "list[dict[str, Any]]":
    rows = [{**row, "kinds": "·".join(_KIND_LABEL.get(k, k) for k in row.get("kinds", []))} for row in summary]
    total = {key: sum(float(row.get(key) or 0) for row in summary) for key in ("pretax", "income_tax", "resident_tax", "other_deduct", "payment", "card_limit")}
    rows.append({"name": "합계", **total, "_bold": True})
    rows.append({})
    rows.append({"name": "전표", "kinds": "계정", "pretax": "차변", "income_tax": "대변", "resident_tax": "적요", "_bold": True})
    for line in journal_lines(summary):
        debit, credit = (line["amount"], None) if line["side"] == "차변" else (None, line["amount"])
        rows.append({"kinds": line["account"], "pretax": debit, "income_tax": credit, "resident_tax": line["remark"]})
    return rows


def build_bonus_workbook(report: "dict[str, Any]") -> bytes:
    workbook = Workbook()
    _sheet(workbook, "주주", SHAREHOLDER_COLUMNS, shareholder_rows(report.get("shareholders", [])), first=True)
    _sheet(workbook, "평·동", ASSOCIATE_COLUMNS, associate_rows(report.get("common", []) + report.get("associates", [])))
    _sheet(workbook, "총괄표", SUMMARY_COLUMNS, summary_rows(report.get("summary", [])))
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
