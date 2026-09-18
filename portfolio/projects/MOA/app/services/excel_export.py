"""목록 화면의 엑셀(xlsx) 내보내기 공통 도우미."""

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

# 한 번에 내보낼 수 있는 최대 행수. 화면 검색 조건이 그대로 적용되므로 보통 수천 건이다.
EXPORT_MAX_ROWS = 100_000


def build_xlsx(
    sheet_name: str, columns: "list[tuple[str, str]]", rows: "list[dict[str, Any]]"
) -> bytes:
    """columns는 (행 dict의 키, 엑셀 헤더) 목록. 순서대로 컬럼이 된다."""
    return build_xlsx_sheets([(sheet_name, columns, rows)])


def build_xlsx_sheets(sheets: "list[tuple[str, list[tuple[str, str]], list[dict[str, Any]]]]") -> bytes:
    """여러 장 — (시트명, columns, rows) 목록 순서대로. 첫 장이 활성 시트다."""
    workbook = Workbook()
    for index, (sheet_name, columns, rows) in enumerate(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        _fill_sheet(sheet, sheet_name, columns, rows)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _fill_sheet(sheet: Any, sheet_name: str, columns: "list[tuple[str, str]]", rows: "list[dict[str, Any]]") -> None:
    sheet.title = sheet_name
    sheet.append([header for _, header in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        values = []
        for key, _ in columns:
            value = row.get(key)
            if isinstance(value, datetime):
                value = value.date()
            elif isinstance(value, Decimal):
                value = float(value)
            values.append(value)
        sheet.append(values)
    for index, (_, header) in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(12, len(header) * 2 + 6)
    sheet.freeze_panes = "A2"


def xlsx_headers(korean_filename: str) -> "dict[str, str]":
    """한글 파일명 다운로드 헤더 (RFC 5987)."""
    stamped = f"{korean_filename}_{date.today():%Y%m%d}.xlsx"
    return {
        "Content-Disposition": (
            f"attachment; filename=export.xlsx; filename*=UTF-8''{quote(stamped)}"
        )
    }


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
