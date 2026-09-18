"""국세청 매출 자료 읽기 (2026-09-11) — 외부 발행 등록에 홈택스 내려받기 파일을 그대로 올린다.

재무팀이 받은 '매출-전자세금계산서 및 현금영수증(국세청 자료).xls' 두 시트:
  세금계산서  '매출 전자(수정) 세금계산서 목록조회' — 6행 머리글(작성일자·승인번호·공급받는자…·품목명…)
  현금영수증  2행 머리글(발행구분·매출일시·공급가액·부가세·봉사료·총금액·승인번호·…·거래구분)
머리글은 위치가 아니라 이름으로 찾는다(몇 행 아래든, 열 순서가 바뀌어도).

이 파일은 외부 발행분만이 아니라 우리 매출 전부다 — 위하고·MOA 팝빌로 끊은 것도 다 들어 있다.
누가 끊었는지는 국세청 승인번호 가운데 8자리가 말해 준다(실측 2026-09-11, 195건):
  41000096 = 위하고(아마란스, bill36524)   41000203 = MOA 팝빌   그 밖(10260909 = 홈택스 직접 등) = 외부
위하고 발급분은 앞으로 이 엑셀로 넣는다 — 읽을 때 출처를 '위하고'로 둔다(사용자 2026-09-11).
MOA 팝빌 발급분은 팝빌 동기화로 들어오므로 검증(external_issue.validate_rows)이 막는다 — 여기서는 읽기만.
"""

import re
from datetime import datetime
from io import BytesIO
from typing import Any

from app.services.external_issue import parse_amount

# 감정서번호(01-2608-3-2642, 변형 -1) 또는 국민약식 번호(400581444) — 품목명·비고에 적힌 것을 뽑는다
DOC_PATTERN = re.compile(r"\d{2}-\d{4}-[0-9A-Z]-\d{4}(?:-\d)?|(?<!\d)4\d{8}(?!\d)")
ORIGINAL_PATTERN = re.compile(r"당초\s*승인번호\s*\(?\s*([0-9A-Za-z-]{20,})")
TAX_HEADERS = ("작성일자", "승인번호", "공급받는자사업자등록번호", "공급가액", "세액")
CASH_HEADERS = ("매출일시", "공급가액", "부가세", "승인번호", "거래구분")
ISSUERS = {"41000096": "위하고", "41000203": "MOA 팝빌"}


def issuer_of(confirm: str) -> str:
    """국세청 승인번호(24자리) → 발급 시스템 이름. 외부 발행이면 빈 문자열."""
    raw = str(confirm or "").replace("-", "").strip()
    return ISSUERS.get(raw[8:16], "") if len(raw) == 24 else ""
def read_sheets(data: bytes, filename: str) -> "list[tuple[str, list[list[Any]]]]":
    """엑셀(.xls BIFF 는 xlrd, .xlsx 는 openpyxl) → [(시트 이름, 행 목록)]. 날짜 셀은 datetime 으로."""
    if filename.lower().endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        sheets = []
        for sheet in book.sheets():
            rows = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                               if cell.ctype == xlrd.XL_CELL_DATE else cell.value)
                rows.append(row)
            sheets.append((sheet.name, rows))
        return sheets
    from openpyxl import load_workbook

    book = load_workbook(BytesIO(data), data_only=True, read_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in book.worksheets]


def _clean(value: Any) -> str:
    return str(value if value is not None else "").replace(" ", "").replace("\n", "").strip()


def _find_header(rows: "list[list[Any]]", needed: "tuple[str, ...]", limit: int = 15) -> "tuple[int, list[str]] | None":
    for index, row in enumerate(rows[:limit]):
        cells = [_clean(v) for v in row]
        if all(name in cells for name in needed):
            return index, cells
    return None


def is_nts(sheets: "list[tuple[str, list[list[Any]]]]") -> bool:
    return any(_find_header(rows, TAX_HEADERS) or _find_header(rows, CASH_HEADERS) for _, rows in sheets)


def _cell(row: "list[Any]", index: "int | None") -> Any:
    return row[index] if index is not None and index < len(row) else None


def _text(value: Any) -> str:
    """화면 표시용 문자열 — 엑셀 날짜 셀은 'YYYY-MM-DD HH:MM:SS'(0시면 날짜만)."""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if (value.hour or value.minute or value.second) else value.strftime("%Y-%m-%d")
    return str(value if value is not None else "").strip()


def _tax_rows(rows: "list[list[Any]]", at: int, cells: "list[str]") -> "list[dict[str, Any]]":
    pos = {name: cells.index(name) for name in TAX_HEADERS}
    corp = pos["공급받는자사업자등록번호"]
    # 상호는 공급자·공급받는자 두 번 나온다 — 공급받는자 사업자번호 뒤의 것
    name_pos = next((i for i in range(corp + 1, len(cells)) if cells[i] == "상호"), None)
    text_pos = [cells.index(h) for h in ("품목명", "비고", "품목비고") if h in cells]
    shown = {name: cells.index(name) for name in ("전자세금계산서분류", "전자세금계산서종류", "발급유형", "품목명", "비고") if name in cells}
    result = []
    for row in rows[at + 1:]:
        confirm = _clean(_cell(row, pos["승인번호"]))
        if not confirm:
            continue
        supply, tax = parse_amount(_cell(row, pos["공급가액"])), parse_amount(_cell(row, pos["세액"]))
        blob = " ".join(str(_cell(row, i) or "") for i in text_pos)
        docs = DOC_PATTERN.findall(blob)
        original = ORIGINAL_PATTERN.search(blob)
        result.append({
            "doc_id": docs[0] if docs else "", "source": "위하고" if issuer_of(confirm) == "위하고" else "국세청",
            # 수정세금계산서(음수)는 취소 — 부호는 종류로 정하고 금액은 양수로 넘긴다(원장 규칙)
            "doc_type": "세금취소" if supply + tax < 0 else "세금계산서",
            "write_date": _cell(row, pos["작성일자"]), "nts_confirm": confirm,
            "receiver_corp_num": _cell(row, corp), "receiver_name": _cell(row, name_pos),
            "supply_cost": abs(supply), "tax": abs(tax),
            "original_confirm": original.group(1).replace("-", "") if original else "",
            # 목록을 국세청 자료 칸 그대로 보여 주려고 원래 값을 들고 간다 (검증·저장엔 안 쓴다)
            "nts": {
                "작성일시": _text(_cell(row, pos["작성일자"])), "승인번호": _text(_cell(row, pos["승인번호"])),
                "분류": _text(_cell(row, shown.get("전자세금계산서분류"))), "종류": _text(_cell(row, shown.get("전자세금계산서종류"))),
                "발급유형": _text(_cell(row, shown.get("발급유형"))), "품목명": _text(_cell(row, shown.get("품목명"))),
                "비고": _text(_cell(row, shown.get("비고"))),
            },
        })
    return result


def _cash_rows(rows: "list[list[Any]]", at: int, cells: "list[str]") -> "list[dict[str, Any]]":
    pos = {name: cells.index(name) for name in CASH_HEADERS}
    shown = {name: cells.index(name) for name in ("발행구분", "비고") if name in cells}
    result = []
    for row in rows[at + 1:]:
        confirm = _clean(_cell(row, pos["승인번호"]))
        if not confirm:
            continue
        supply, vat = parse_amount(_cell(row, pos["공급가액"])), parse_amount(_cell(row, pos["부가세"]))
        cancel = "취소" in _clean(_cell(row, pos["거래구분"])) or supply + vat < 0
        sold_at = _cell(row, pos["매출일시"])
        result.append({
            # 현금영수증 자료에는 감정서번호·거래처가 없다 — 목록에서 감정서번호를 적어 넣는다
            "doc_id": "", "source": "국세청", "doc_type": "현금취소" if cancel else "현금영수증",
            "write_date": sold_at.date() if isinstance(sold_at, datetime) else str(sold_at or "")[:10],
            "nts_confirm": confirm, "receiver_corp_num": "", "receiver_name": "",
            "supply_cost": abs(supply), "tax": abs(vat), "original_confirm": "",
            "nts": {
                "작성일시": _text(sold_at), "승인번호": confirm, "분류": "현금영수증",
                "종류": _text(_cell(row, pos["거래구분"])), "발급유형": _text(_cell(row, shown.get("발행구분"))),
                "품목명": "", "비고": _text(_cell(row, shown.get("비고"))),
            },
        })
    return result


def parse_nts(sheets: "list[tuple[str, list[list[Any]]]]") -> "list[dict[str, Any]]":
    """세금계산서 시트 → 현금영수증 시트 순으로 행을 모은다(검증 전 원자료)."""
    rows: "list[dict[str, Any]]" = []
    for _, sheet_rows in sheets:
        tax = _find_header(sheet_rows, TAX_HEADERS)
        if tax:
            rows.extend(_tax_rows(sheet_rows, *tax))
            continue
        cash = _find_header(sheet_rows, CASH_HEADERS)
        if cash:
            rows.extend(_cash_rows(sheet_rows, *cash))
    return rows
