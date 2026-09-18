"""협회 업무실적 '데이터' 시트를 양식(.xlsx)에 채워 내보낸다.

app/static/templates/work_report_base.xlsx (데이터 헤더 2행 + 협회코드 참조시트)를
베이스로 열고, build_kapa_rows()가 만든 행을 3행부터 협회 컬럼 순서로 기록한다.
"""

from datetime import date, datetime
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

_BASE = Path(__file__).parents[1] / "static" / "templates" / "work_report_base.xlsx"
_DATA_SHEET = "데이터"
_FIRST_ROW = 3  # 헤더 2행 다음부터

# (엑셀 컬럼 0-index, 매핑 키, 종류). 종류: text/int/num/date. 빈 항목(자동입력x·평가사2 등)은
# 협회 프로그램이 채우므로 비워 둔다.
_COLUMNS: list[tuple[int, str, str]] = [
    (0, "BUNGI", "text"), (1, "MON", "int"), (2, "ID_NUM", "text"),
    (3, "QUANO", "text"), (4, "GAMMAN", "text"),
    (7, "PCODE", "text"), (9, "YCODE", "text"),
    (11, "GNAME", "text"), (12, "CUST", "text"), (13, "CUSTCODE", "text"),
    (15, "CONSULTDATE", "date"), (16, "IN_DATE", "date"), (17, "OUTDATE", "date"),
    (18, "GAMGA", "num"), (19, "FEE", "num"), (20, "SUSU", "num"),
    (21, "REG", "text"), (22, "EUB", "text"), (23, "SAN", "text"),
    (24, "BUN1", "text"), (25, "BUN2", "text"),
]
# 물건구분 1~9: 협회코드(CATEGORY)·물건수(CNT)·평가액(PRICE) 3칸씩 (양식 AA=26부터).
for _i in range(1, 10):
    _base = 26 + 3 * (_i - 1)
    _COLUMNS.append((_base, f"CATEGORY{_i}", "text"))
    _COLUMNS.append((_base + 1, f"CNT{_i}", "int"))
    _COLUMNS.append((_base + 2, f"PRICE{_i}", "num"))


@lru_cache
def _base_bytes() -> bytes:
    if not _BASE.exists():
        raise FileNotFoundError(f"양식 베이스가 없습니다: {_BASE}")
    return _BASE.read_bytes()


def _cell_value(kind: str, raw: Any) -> Any:
    if raw is None or raw == "":
        return None
    if kind == "date":
        text = raw.isoformat() if isinstance(raw, (datetime, date)) else str(raw)
        return text[:10]  # 'YYYY-MM-DD' (양식은 날짜를 텍스트로 보관)
    if kind == "num":
        try:
            return float(raw)
        except (TypeError, ValueError):
            return None
    if kind == "int":
        try:
            return int(float(raw))
        except (TypeError, ValueError):
            return None
    value = str(raw).strip()
    return value or None


def build_work_report_xlsx(result: dict[str, Any]) -> bytes:
    """build_kapa_rows() 결과를 협회 양식(.xlsx) 바이트로. 각 행에 분기·월도 채운다."""
    workbook = load_workbook(BytesIO(_base_bytes()))
    sheet = workbook[_DATA_SHEET]
    shared = {"BUNGI": result.get("bungi"), "MON": result.get("month")}
    for offset, row in enumerate(result.get("rows", [])):
        excel_row = _FIRST_ROW + offset
        for col_index, key, kind in _COLUMNS:
            raw = shared.get(key, row.get(key))
            value = _cell_value(kind, raw)
            if value is not None:
                sheet.cell(row=excel_row, column=col_index + 1, value=value)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
