"""업무실적 엑셀에서 감정서번호를 추출한다."""

import io
import re
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

DOC_RE = re.compile(r"\d{2}-\d{4}-[A-Za-z0-9]-\d{4}")
MAX_ROWS = 100_000


def _extract_docs(rows: Iterable[Iterable[Any]]) -> list[str]:
    docs: list[str] = []
    seen: set[str] = set()
    for row_index, row in enumerate(rows):
        if row_index >= MAX_ROWS:
            break
        for value in row:
            if value is None:
                continue
            for match in DOC_RE.finditer(str(value)):
                doc = match.group(0).upper()
                if doc not in seen:
                    seen.add(doc)
                    docs.append(doc)
    return docs


def extract_work_report_doc_ids(data: bytes, filename: str) -> list[str]:
    """.xlsx/.xls의 모든 시트에서 감정서번호를 순서대로 중복 없이 반환한다."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        try:
            docs: list[str] = []
            seen: set[str] = set()
            for sheet in workbook.worksheets:
                for doc in _extract_docs(sheet.iter_rows(values_only=True)):
                    if doc not in seen:
                        seen.add(doc)
                        docs.append(doc)
            return docs
        finally:
            workbook.close()
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        try:
            docs = []
            seen = set()
            for sheet in workbook.sheets():
                for doc in _extract_docs(
                    (sheet.row_values(index) for index in range(sheet.nrows))
                ):
                    if doc not in seen:
                        seen.add(doc)
                        docs.append(doc)
            return docs
        finally:
            workbook.release_resources()
    raise ValueError("엑셀 파일(.xls, .xlsx)만 올릴 수 있습니다.")
