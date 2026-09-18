"""엑셀 지급 시트 읽기·대조 (2단계 15) — 가짜 시트로 블록 주인 귀속과 합계 읽기를 본다."""

from openpyxl import Workbook

from app.services.bonus.excel_compare import compare, read_associate_sheet, read_shareholder_sheet


def _workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "26.08월-주주"
    # r7 상세(김형식 D열이지만 블록 주인은 강무진), r8 합계(강무진 40%), r9 상세, r10 합계(김형식 30%)
    ws.append([])
    for _ in range(5):
        ws.append([])
    ws.cell(7, 1, "01-2503-5-0042"); ws.cell(7, 4, "김형식"); ws.cell(7, 6, 40_000_000); ws.cell(7, 8, 10_000_000)
    ws.cell(8, 1, "주주 합계"); ws.cell(8, 4, "강무진"); ws.cell(8, 17, 9_900_000); ws.cell(8, 18, 3_960_000)
    ws.cell(8, 25, 3_960_000); ws.cell(8, 26, 1_188_000); ws.cell(8, 27, 118_800); ws.cell(8, 29, 2_653_200); ws.cell(8, 31, 0)
    ws.cell(9, 1, "01-2503-5-0042"); ws.cell(9, 4, "김형식"); ws.cell(9, 6, 40_000_000); ws.cell(9, 8, 20_000_000)
    ws.cell(10, 1, "주주 합계"); ws.cell(10, 4, "김형식"); ws.cell(10, 17, -100); ws.cell(10, 21, 0); ws.cell(10, 25, 0); ws.cell(10, 31, -100)
    pd = wb.create_sheet("26.08월-평.동")
    for _ in range(7):
        pd.append([])
    pd.cell(8, 1, "01-2606-6-0406"); pd.cell(8, 4, "공(김기석)"); pd.cell(8, 7, 10_000); pd.cell(8, 12, 9_900); pd.cell(8, 27, 297)
    pd.cell(9, 3, "주주 합계"); pd.cell(9, 4, "김기석"); pd.cell(9, 27, 297); pd.cell(9, 30, 0); pd.cell(9, 31, 0); pd.cell(9, 32, 0); pd.cell(9, 34, 0)
    pd.cell(10, 1, "01-2508-4-0279"); pd.cell(10, 4, "공(이영은)"); pd.cell(10, 7, 9_166_600); pd.cell(10, 12, 9_074_934); pd.cell(10, 27, 907_493.4)
    pd.cell(11, 3, "소속 합계"); pd.cell(11, 4, "이영은"); pd.cell(11, 27, 907_493.4); pd.cell(11, 30, 634_900); pd.cell(11, 31, 190_000); pd.cell(11, 32, 19_000); pd.cell(11, 34, 425_900)
    return wb


def test_주주_시트는_블록_주인에게_귀속하고_합계를_사람별로_더한다():
    sheet = read_shareholder_sheet(_workbook()["26.08월-주주"])
    assert sheet["rows"][("01-2503-5-0042", "강무진")] == {"fee": 40_000_000, "assessed": 10_000_000}
    assert sheet["rows"][("01-2503-5-0042", "김형식")] == {"fee": 40_000_000, "assessed": 20_000_000}
    kang = sheet["totals"]["강무진"]
    assert (kang["payout_base"], kang["bonus"], kang["pretax"], kang["payment"], kang["blocks"]) == (9_900_000, 3_960_000, 3_960_000, 2_653_200, 1)
    assert sheet["totals"]["김형식"]["unpaid_carry_out"] == -100


def test_평동_시트는_공통건_주주와_소속을_구분한다():
    sheet = read_associate_sheet(_workbook()["26.08월-평.동"])
    assert sheet["rows"][("01-2606-6-0406", "김기석")] == {"fee": 10_000, "assessed": 9_900, "bonus": 297}
    assert sheet["totals"]["김기석"]["kind"] == "COMMON"
    assert sheet["totals"]["이영은"] == {"kind": "ASSOCIATE", "bonus": 907_493.4, "pretax": 634_900, "income_tax": 190_000, "resident_tax": 19_000, "other_deduct": 0, "payment": 425_900}


def test_대조는_감정서_커버리지와_사람별_차이를_낸다():
    wb = _workbook()
    report = {
        "shareholders": [{"name": "강무진", "rows": [{"doc_id": "01-2503-5-0042", "fee": 10_000_000, "assessed": 10_000_000}],
                          "totals": {"payout_base": 9_900_000, "bonus": 3_960_000, "pretax": 3_960_000, "income_tax": 1_188_000, "resident_tax": 118_800, "payment": 2_653_200, "unpaid_carry_out": 0}}],
        "common": [{"name": "김기석", "rows": [{"doc_id": "01-2606-6-0406", "fee": 10_000, "assessed": 9_900}],
                    "totals": {"bonus": 297, "pretax": 0, "income_tax": 0, "resident_tax": 0, "payment": 0}}],
        "associates": [],
        "held": [{"doc_id": "01-2508-4-0279", "reason": "HELD_UNPAID"}],
    }
    result = compare(report, read_shareholder_sheet(wb["26.08월-주주"]), read_associate_sheet(wb["26.08월-평.동"]))
    assert result["docs"]["excel"] == 3 and result["docs"]["engine"] == 2 and result["docs"]["both"] == 2
    assert result["docs"]["excel_only"] == ["01-2508-4-0279"] and result["docs"]["held_in_excel"] == ["01-2508-4-0279"]
    assert result["pairs"] == {"excel": 4, "engine": 2, "both": 2, "close": 2}
    by_name = {p["name"]: p for p in result["persons"]}
    assert by_name["강무진"]["payment_exact"] is True and by_name["강무진"]["pretax_close"] is True
    assert by_name["이영은"]["engine"]["payment"] == 0 and by_name["이영은"]["gap"] == -634_900
    assert result["persons"][0]["name"] == "이영은"        # 차이 큰 순
