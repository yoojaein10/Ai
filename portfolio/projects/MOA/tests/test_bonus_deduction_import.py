"""공제 대장 엑셀 올리기 — 양식 파싱(한글 종류·적용월·오류 행)과 양식 파일."""

from datetime import date
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.services.bonus import deduction_import as di


def _book(rows, header=("사람", "종류", "금액", "감정서", "발생일", "메모", "적용월")) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_한글_종류와_적용월을_읽고_오류_행은_건너뛴다():
    content = _book([
        ("강무진", "화환", 80000, None, "2026-08-12", "부친상", None),
        ("김정원", "감정서경비", "20,000", "01-2606-5-0085", date(2026, 8, 15), "전자수입인지", "2026-09"),
        ("고세욱", "ADVANCE_PAID", 4610400, "01-2505-4-0176", None, None, 202609),
        ("", "화환", 1, None, None, None, None),                                # 사람 없음
        ("박용준", "이상한종류", 100, None, None, None, None),                  # 종류 모름
        ("박용준", "화환", 0, None, None, None, None),                         # 금액 0
        ("박용준", "화환", 100, None, None, None, "26.9"),                     # 적용월 형식 오류
        (None, None, None, None, None, None, None),                          # 빈 줄
    ])
    rows, errors = di.parse_workbook(content, source_label="xlsx")
    assert [(r["person"], r["kind"], r["amount"], r["doc_id"], r["occurred_on"], r["apply_period"]) for r in rows] == [
        ("강무진", "WREATH", 80000.0, None, date(2026, 8, 12), None),
        ("김정원", "DOC_EXPENSE", 20000.0, "01-2606-5-0085", date(2026, 8, 15), "202609"),
        ("고세욱", "ADVANCE_PAID", 4610400.0, "01-2505-4-0176", None, "202609"),
    ]
    assert rows[0]["source"] == "MANUAL" and rows[0]["source_key"].startswith("xlsx:") and rows[0]["source_key"].endswith(":2")
    assert [e.split(":")[0] for e in errors] == ["5행", "6행", "7행", "8행"]
    assert "종류를 모름" in errors[1] and "적용월" in errors[3]


def test_머리글이_없으면_전체가_오류다():
    rows, errors = di.parse_workbook(_book([("강무진", 1)], header=("이름", "값")))
    assert rows == [] and "종류" in errors[0] and "금액" in errors[0]


def test_양식_파일은_머리글과_종류_안내를_담는다():
    wb = load_workbook(BytesIO(di.template_workbook()))
    assert wb.sheetnames == ["공제", "종류"]
    assert [c.value for c in wb["공제"][1]] == ["사람", "종류", "금액", "감정서", "발생일", "메모", "적용월"]
    kinds = {row[1].value for row in wb["종류"].iter_rows(min_row=2)}
    assert kinds == set(di.KIND_LABELS.values())
