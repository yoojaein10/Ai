"""엑셀 공제내역 두 시트 → 공제 대장. 월 블록은 월 시트의 O·W 수식이 가리키는 행 범위로 정한다."""

import sqlite3
from datetime import date, datetime

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import BonusClose, BonusDeduction
from app.services.bonus import deduction_items, excel_deductions
from app.services.bonus.ledger import load_deductions

sqlite3.register_adapter(date, lambda value: value.isoformat())

DOC_SHEET = "감정서관련비용공제내역"
PERSON_SHEET = "화환및감정서관련없는비용공제내역"
# 주주 시트 합계행: (행, 이름, W 값, Y 수식 꼬리(선지급), AB 수식)
OWNERS = ((12, "강무진", 80_000, "-20000000", "=27600+26000"), (20, "권오억", 0, "", None), (24, "송정선", 0, "", None), (28, "김형식", 0, "", None))


def _month_sheets(wb, *, values: bool):
    ws = wb.create_sheet("26.08월-주주")
    for r, owner, wreath, advance, other in OWNERS:
        ws.cell(r, 1, "주주 합계")
        ws.cell(r, 4, owner)
        if values:
            ws.cell(r, 23, wreath)
            ws.cell(r, 28, 53_600 if other else None)
        else:
            ws.cell(r, 15, f"=SUMIFS({DOC_SHEET}!$E$5:$E$9,{DOC_SHEET}!$D$5:$D$9,'26.08월-주주'!D{r})")
            ws.cell(r, 23, f"=IF(Q{r}<0,0,SUMIFS({PERSON_SHEET}!$D$3:$D$7,{PERSON_SHEET}!$B$3:$B$7,'26.08월-주주'!D{r}))")
            ws.cell(r, 25, f"=ROUNDDOWN(R{r}+V{r}-W{r}+X{r},-3){advance}")
            ws.cell(r, 28, other)
    pd = wb.create_sheet("26.08월-평.동")                       # 소속평가사 김기도 — 화환은 여기 AC 에서 빠진다
    pd.cell(49, 3, "소속 합계"); pd.cell(49, 4, "김기도"); pd.cell(49, 29, 400_000 if values else "=IF(AA49<0,0,SUMIFS(...))")


def _doc_sheet(wb):
    ws = wb.create_sheet(DOC_SHEET)
    ws.append(["날짜", "담당자", "내역", "금액"])
    ws.cell(5, 1, datetime(2026, 7, 15)); ws.cell(5, 3, "01-2604-1-0239 전자수입인지"); ws.cell(5, 4, "강무진"); ws.cell(5, 5, 20_000); ws.cell(5, 7, "세금과공과금")
    ws.cell(6, 1, "2026-07-31"); ws.cell(6, 3, "7월 제본 01-2603-1-0011"); ws.cell(6, 4, "권오억"); ws.cell(6, 5, 8_182); ws.cell(6, 7, "도서인쇄비")
    ws.cell(7, 1, datetime(2026, 7, 16)); ws.cell(7, 3, "우신빌라 재건축 전자수입인지"); ws.cell(7, 4, "김형식"); ws.cell(7, 5, 20_000); ws.cell(7, 7, "세금과공과금")
    ws.cell(8, 1, datetime(2026, 7, 15)); ws.cell(8, 3, "화곡초일대 전자수입인지"); ws.cell(8, 4, "정인수"); ws.cell(8, 5, 40_000); ws.cell(8, 7, "세금과공과금")   # 주주 시트에 없는 사람
    ws.cell(9, 1, datetime(2026, 7, 20)); ws.cell(9, 3, "01-2605-5-0076 잡급 환입"); ws.cell(9, 4, "송정선"); ws.cell(9, 5, -1_500_000); ws.cell(9, 7, "잡급")
    ws.cell(10, 5, 48_182)                                                                   # 소계 행 — 담당자 없음
    return ws


def _person_sheet(wb):
    ws = wb.create_sheet(PERSON_SHEET)
    ws.append(["날짜", "담당자", "내역", "금액", "", "공제월"])
    ws.cell(3, 1, datetime(2026, 7, 9)); ws.cell(3, 2, "강무진"); ws.cell(3, 4, 80_000); ws.cell(3, 5, "화환")
    ws.cell(4, 1, "26.03"); ws.cell(4, 2, "권오억"); ws.cell(4, 3, "01-2603-3-0001"); ws.cell(4, 4, 100_000); ws.cell(4, 5, "신한은행 패널티")
    ws.cell(5, 1, datetime(2026, 7, 7)); ws.cell(5, 2, "권오억"); ws.cell(5, 4, 377_290); ws.cell(5, 5, "법인차량")
    ws.cell(6, 1, datetime(2026, 7, 31)); ws.cell(6, 2, "강무진"); ws.cell(6, 3, "신문대금"); ws.cell(6, 4, 20_000); ws.cell(6, 5, "도서인쇄비")
    ws.cell(7, 1, "26.03"); ws.cell(7, 2, "김기도"); ws.cell(7, 4, 400_000); ws.cell(7, 5, "신한은행 패널티")        # 평·동 AC 로 적용
    ws.cell(8, 4, 977_290)                                                                   # 소계
    ws.cell(10, 1, "공제못함")
    ws.cell(11, 1, datetime(2026, 7, 28)); ws.cell(11, 2, "유승연"); ws.cell(11, 4, 100_000); ws.cell(11, 5, "화환")
    ws.cell(12, 1, datetime(2026, 7, 9)); ws.cell(12, 2, "강무진"); ws.cell(12, 4, 80_000); ws.cell(12, 5, "화환")   # r3 과 같은 항목이 뒤에 또 (복사) → 건너뜀
    return ws


def _workbooks():
    wb_f, wb_v = Workbook(), Workbook()
    for wb, values in ((wb_f, False), (wb_v, True)):
        wb.remove(wb.active)
        _month_sheets(wb, values=values)
        _doc_sheet(wb)
        _person_sheet(wb)
    return wb_f, wb_v


def test_월_블록은_월_시트_수식의_행_범위다():
    wb_f, _ = _workbooks()
    assert excel_deductions.read_month_blocks(wb_f) == {"202608": {"doc": (5, 9), "person": (3, 7)}}


def test_감정서_관련_공제는_감정서번호를_뽑아_그_달에_적용되고_시트에_없는_사람은_대기다():
    wb_f, wb_v = _workbooks()
    items = [i for i in excel_deductions.build_items(wb_f, wb_v) if i["kind"] in ("DOC_EXPENSE", "EXPENSE_CREDIT")]
    assert [(i["person"], i["kind"], i["doc_id"], i["amount"], i["occurred_on"], i["status"], i["applied_period"]) for i in items] == [
        ("강무진", "DOC_EXPENSE", "01-2604-1-0239", 20_000, date(2026, 7, 15), "APPLIED", "202608"),
        ("권오억", "DOC_EXPENSE", "01-2603-1-0011", 8_182, date(2026, 7, 31), "APPLIED", "202608"),
        ("김형식", "DOC_EXPENSE", None, 20_000, date(2026, 7, 16), "APPLIED", "202608"),
        ("정인수", "DOC_EXPENSE", None, 40_000, date(2026, 7, 15), "PENDING", None),
        ("송정선", "EXPENSE_CREDIT", "01-2605-5-0076", 1_500_000, date(2026, 7, 20), "APPLIED", "202608"),
    ]
    assert items[0]["memo"] == "세금과공과금 01-2604-1-0239 전자수입인지" and items[0]["source_key"] == "doc:26.08:5"
    assert "주주 시트에 없음" in items[3]["memo"]


def test_화환류는_W가_0이면_대기로_남고_평동_AC로도_적용되며_블록_밖과_복사본은_대기이거나_건너뛴다():
    wb_f, wb_v = _workbooks()
    items = [i for i in excel_deductions.build_items(wb_f, wb_v) if i["kind"] in deduction_items.CARRY_KINDS]
    by = {(i["person"], i["amount"]): i for i in items}
    assert by[("강무진", 80_000)]["status"] == "APPLIED" and by[("강무진", 80_000)]["kind"] == "WREATH"
    assert by[("강무진", 20_000)]["kind"] == "OTHER_EXPENSE" and by[("강무진", 20_000)]["memo"] == "도서인쇄비 신문대금"
    assert by[("권오억", 100_000)]["kind"] == "PENALTY" and by[("권오억", 100_000)]["status"] == "PENDING"
    assert "Q<0" in by[("권오억", 100_000)]["memo"] and by[("권오억", 100_000)]["doc_id"] == "01-2603-3-0001"
    assert by[("권오억", 377_290)]["kind"] == "INSURANCE" and by[("권오억", 377_290)]["status"] == "PENDING"
    assert by[("김기도", 400_000)]["status"] == "APPLIED" and by[("김기도", 400_000)]["applied_period"] == "202608"
    assert by[("유승연", 100_000)]["status"] == "PENDING" and by[("유승연", 100_000)]["applied_period"] is None
    assert len([i for i in items if i["person"] == "강무진" and i["amount"] == 80_000]) == 1           # r12 복사본은 건너뜀


def test_선지급과_기타공제는_합계행에서_읽는다():
    wb_f, wb_v = _workbooks()
    items = excel_deductions.build_items(wb_f, wb_v)
    advance = next(i for i in items if i["kind"] == "ADVANCE_PAID")
    assert (advance["person"], advance["amount"], advance["applied_period"], advance["source_key"]) == ("강무진", 20_000_000, "202608", "adv:26.08:12")
    other = next(i for i in items if i["kind"] == "OTHER_DEDUCT")
    assert (other["person"], other["amount"], other["memo"]) == ("강무진", 53_600, "엑셀 AB =27600+26000")


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[BonusClose.__table__, BonusDeduction.__table__])
    session = sessionmaker(bind=engine, future=True)()
    try:
        yield session
    finally:
        session.close()


def test_시딩은_멱등이고_마감된_달에도_이력으로_넣는다(db):
    db.add(BonusClose(period="202608", status="CLOSED", source="EXCEL")); db.commit()
    wb_f, wb_v = _workbooks()
    preview = excel_deductions.seed_deductions(db, None, workbooks=(wb_f, wb_v), dry_run=True)
    assert preview["dry_run"] is True and preview["written"] == 13 and db.query(BonusDeduction).count() == 0
    report = excel_deductions.seed_deductions(db, None, workbooks=(wb_f, wb_v))
    assert report["written"] == 13 and report["skipped"] == 0            # 감정서 5 + 화환류 5 + 선지급 1 + AB 1 + 공제못함 1
    assert report["by_period"] == {"202608": 9, "PENDING": 4}
    applied = load_deductions(db, "202608")
    assert sorted(applied) == ["강무진", "권오억", "김기도", "김형식", "송정선"]
    assert sorted((i["kind"], i["amount"]) for i in applied["강무진"]) == [("ADVANCE_PAID", 20_000_000), ("DOC_EXPENSE", 20_000), ("OTHER_DEDUCT", 53_600), ("OTHER_EXPENSE", 20_000), ("WREATH", 80_000)]
    assert report["pending_people"] == {"권오억": 477_290, "정인수": 40_000, "유승연": 100_000}
    again = excel_deductions.seed_deductions(db, None, workbooks=(wb_f, wb_v))
    assert again["written"] == 0 and again["skipped"] == 13 and db.query(BonusDeduction).count() == 13
    forced = excel_deductions.seed_deductions(db, None, workbooks=(wb_f, wb_v), force=True)
    assert forced["written"] == 13 and db.query(BonusDeduction).count() == 13
