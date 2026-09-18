"""엑셀 지급 이력 시딩 (2단계 16) — 도입 전 달을 EXCEL 마감으로 넣어야 두 번 지급하지 않고 첫 달 이월이 나온다."""

import sqlite3
from datetime import date

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import BonusClose, BonusResult
from app.services.bonus import closing
from app.services.bonus.excel_history import seed_history

sqlite3.register_adapter(date, lambda value: value.isoformat())


def _workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "26.07월-주주"
    for _ in range(6):
        ws.append([])
    ws.cell(7, 1, "01-2503-5-0042"); ws.cell(7, 4, "김형식"); ws.cell(7, 6, 40_000_000); ws.cell(7, 8, 10_000_000)
    ws.cell(8, 1, "주주 합계"); ws.cell(8, 3, "2019년분/2021년 3월~"); ws.cell(8, 4, "강무진")
    ws.cell(8, 17, 9_900_000); ws.cell(8, 18, 3_960_000); ws.cell(8, 25, 3_960_000); ws.cell(8, 26, 1_188_000)
    ws.cell(8, 27, 118_800); ws.cell(8, 29, 2_653_200); ws.cell(8, 31, 0)
    ws.cell(9, 1, "주주 합계"); ws.cell(9, 3, "2020년~2021년 2월"); ws.cell(9, 4, "강무진"); ws.cell(9, 17, 0); ws.cell(9, 25, 0)
    ws.cell(10, 1, "주주 합계"); ws.cell(10, 3, "2019년분/2021년 3월~"); ws.cell(10, 4, "권오억")
    ws.cell(10, 17, -2_693_660); ws.cell(10, 18, 0); ws.cell(10, 25, 0); ws.cell(10, 29, 0); ws.cell(10, 31, -2_693_660)
    pd = wb.create_sheet("26.07월-평.동")
    for _ in range(7):
        pd.append([])
    pd.cell(8, 1, "01-2606-6-0406"); pd.cell(8, 4, "공(김기석)"); pd.cell(8, 7, 10_000); pd.cell(8, 12, 9_900); pd.cell(8, 27, 297)
    pd.cell(9, 3, "주주 합계"); pd.cell(9, 4, "김기석"); pd.cell(9, 27, 297); pd.cell(9, 30, 0); pd.cell(9, 31, 0); pd.cell(9, 32, 0); pd.cell(9, 34, 0)
    pd.cell(10, 1, "01-2508-4-0279"); pd.cell(10, 4, "공(이영은)"); pd.cell(10, 7, 9_166_600); pd.cell(10, 12, 9_074_934); pd.cell(10, 27, 907_493.4)
    pd.cell(11, 3, "소속 합계"); pd.cell(11, 4, "이영은"); pd.cell(11, 27, 907_493.4); pd.cell(11, 30, 634_900); pd.cell(11, 31, 190_000); pd.cell(11, 32, 19_000); pd.cell(11, 34, 425_900)
    wb.create_sheet("총괄표 07")
    wb.save(path)
    return path


@pytest.fixture()
def workbook(tmp_path):
    return _workbook(tmp_path / "성과상여.xlsx")


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[BonusClose.__table__, BonusResult.__table__])
    session = sessionmaker(bind=engine, future=True)()
    try:
        yield session
    finally:
        session.close()


def test_월_시트가_EXCEL_마감과_결과_행이_된다(workbook, db):
    report = seed_history(db, workbook)
    assert report["written"] == ["202607"] and report["skipped"] == []
    assert closing.close_status(db, "202607") == "CLOSED" and closing.close_source(db, "202607") == "EXCEL"
    rows = closing.load_result(db, "202607")
    docs = {(r["doc_id"], r["person"]): r for r in rows if r["doc_id"]}
    assert float(docs[("01-2503-5-0042", "강무진")]["fee"]) == 10_000_000       # F × (H/(F+G))
    assert float(docs[("01-2503-5-0042", "강무진")]["assessed"]) == 10_000_000
    assert docs[("01-2606-6-0406", "김기석")]["kind"] == "COMMON"
    assert docs[("01-2508-4-0279", "이영은")]["kind"] == "ASSOCIATE"
    totals = {(r["person"], r["kind"]): r for r in rows if r["doc_id"] is None}
    kang = totals[("강무진", "SHAREHOLDER")]
    assert (float(kang["payout_base"]), float(kang["bonus"]), float(kang["pretax"]), float(kang["payment"])) == (9_900_000, 3_960_000, 3_960_000, 2_653_200)
    assert float(totals[("권오억", "SHAREHOLDER")]["unpaid_carry_out"]) == -2_693_660
    assert float(totals[("이영은", "ASSOCIATE")]["payment"]) == 425_900
    assert all(r["source"] == "EXCEL" for r in rows)


def test_이력이_있으면_다음_달_이월과_기지급이_나온다(workbook, db):
    seed_history(db, workbook)
    assert closing.carry_in(db, "202608") == {"강무진": 0.0, "권오억": 2_693_660.0, "김기석": 0.0, "이영은": 0.0}
    assert closing.already_paid(db, ["01-2503-5-0042"]) == {("01-2503-5-0042", "강무진"): 10_000_000}


def test_다시_돌리면_건너뛰고_force_로만_다시_쓴다(workbook, db):
    seed_history(db, workbook)
    again = seed_history(db, workbook)
    assert again["written"] == [] and again["skipped"] == ["202607"]
    assert len(closing.load_result(db, "202607")) == 7
    forced = seed_history(db, workbook, force=True)
    assert forced["written"] == ["202607"]
    assert len(closing.load_result(db, "202607")) == 7                 # 두 번 쌓이지 않는다
    dry = seed_history(db, workbook, periods=["202607"], dry_run=True, force=True)
    assert dry["dry_run"] is True and dry["written"] == ["202607"]


def test_MOA_로_마감한_달은_건드리지_않는다(workbook, db):
    db.add(BonusClose(period="202607", status="CLOSED", source="MOA"))
    db.add(BonusResult(period="202607", person="강무진", kind="SHAREHOLDER", doc_id=None, fee=0, source="MOA"))
    db.commit()
    report = seed_history(db, workbook, force=True)
    assert report["written"] == [] and report["protected"] == ["202607"]
    assert len(closing.load_result(db, "202607")) == 1
