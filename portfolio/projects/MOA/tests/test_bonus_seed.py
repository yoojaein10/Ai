"""엑셀 → 마스터 시드 (1단계 5). 가짜 워크북을 만들어 읽기·저장을 검증한다.

실제 파일은 시트가 54개고 수식이 있는 달(26.02~)과 값만 있는 달(25.12·26.01)이 섞여
있다. 시딩은 (1) 요율 = 가장 최근 시트의 사람별 블록 라벨, (2) 지분 = 상세행을 다음
합계행의 주인에게 귀속해 수식 %(없으면 H/(F+G) 값), (3) 소속 파라미터 = '소속 합계'
행의 AD(×70%)·AE(×15%/30%) 수식에서 읽는다.
"""

import sqlite3
from datetime import date

import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import BonusPerson, BonusRate, BonusShare
from app.services.bonus import excel_seed, persons, schedule, shares
from scripts.seed_bonus_master import seed

sqlite3.register_adapter(date, lambda value: value.isoformat())


def _shareholder_sheet(wb, title, rows):
    ws = wb.create_sheet(title)
    ws["A4"], ws["C4"], ws["D4"], ws["F4"], ws["H4"] = "감정서번호", "거래처", "담당자", "순수수료", "산정금액"
    ws["R6"], ws["S6"], ws["T6"], ws["U6"] = 0.4, 0.45, 0.35, 0.3
    r = 7
    for row in rows:
        if row[0] == "합계":
            _, label, owner, rate_col = row
            ws.cell(r, 1, "주주 합계"); ws.cell(r, 3, label); ws.cell(r, 4, owner)
            col = {40: 18, 45: 19, 35: 20, 30: 21}[rate_col]
            ws.cell(r, col, f"=IF(Q{r}<0,0,Q{r}*${chr(64 + col)}$6)")
        else:
            doc, work, customer, manager, fee, land, assessed = row
            ws.cell(r, 1, doc); ws.cell(r, 2, work); ws.cell(r, 3, customer); ws.cell(r, 4, manager)
            ws.cell(r, 6, fee); ws.cell(r, 7, land)
            ws.cell(r, 8, assessed.replace("{r}", str(r)) if isinstance(assessed, str) else assessed)
        r += 1
    return ws


def _associate_sheet(wb, title, totals):
    ws = wb.create_sheet(title)
    ws["A5"], ws["C5"], ws["D5"] = "감정서번호", "거래처", "담당자"
    r = 8
    for person, ad_formula, ae_formula in totals:
        ws.cell(r, 3, "소속 합계"); ws.cell(r, 4, person)
        ws.cell(r, 30, ad_formula.replace("{r}", str(r))); ws.cell(r, 31, ae_formula.replace("{r}", str(r)))
        r += 1
    return ws


@pytest.fixture()
def workbook(tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("총괄표 08")  # 월 시트가 아닌 것은 무시해야 한다
    _shareholder_sheet(wb, "26.07월-주주", [
        ("01-2605-3-1564", "담보", "농협 어딘가", "노승환", 500000, None, "=(F{r}+G{r})"),
        ("합계", "2024년 7월~", "노승환", 30),
    ])
    _shareholder_sheet(wb, "26.08월-주주", [
        ("01-2503-5-0042", "컨설팅", "신림제7구역 (김5 강2.5 조2.5)", "김형식", 40000000, None, "=(F{r}+G{r})*25%"),
        ("합계", "2019년분/2021년 3월~", "강무진", 40),
        ("합계", "2020년~2021년 2월", "강무진", 45),
        ("01-2603-1-0155", "국공유재산", "주택도시보증공사 든든전세임대 (안9:이1)", "안창덕", 284000, None, 255600),
        ("01-2607-3-0001", "담보", "이지안 (HF Web 보증) (안6:황4/법카 안100)", "안창덕", 1000000, 0, "=(F{r}+G{r})*60%"),
        ("합계", "2019년분/2021년 3월~", "안창덕", 40),
        ("01-2607-3-0002", "담보", "빈 건", "김형식", 0, 0, None),
        ("01-2503-5-0042", "컨설팅", "신림제7구역 (김5 강2.5 조2.5)", "김형식", 40000000, None, "=(F{r}+G{r})*50%"),
        ("01-2606-3-1813", "담보", "혼자 하는 건", "김형식", 913814, None, "=(F{r}+G{r})"),
        ("합계", "2022년 7월~", "김형식", 30),
        ("01-2606-3-1930", "담보", "농협 어딘가", "노승환", 700000, None, "=(F{r}+G{r})"),
        ("합계", "2025년 7월 7일 접수분~", "노승환", 40),
        ("합계", "2022년 7월~ 2025년 7월 6일 접수분까지", "노승환", 30),
        ("01-2607-6-0999", "가격자문", "어느 은행", "황인석", 45000, None, "=(F{r}+G{r})"),
        ("합계", "이상한 라벨", "황인석", 40),
        # 이영은: 주주였다가(~2023.12) 지금은 소속 — 같은 달 두 시트에 다 있다
        ("01-2306-3-0001", "담보", "옛날 건", "이영은", 300000, None, "=(F{r}+G{r})"),
        ("합계", "2019년분/2021년 3월~2023.12", "이영은", 40),
    ])
    _associate_sheet(wb, "26.08월-평.동", [
        ("이영은", "=ROUNDDOWN(AA{r}+AB{r}-AC{r},-3)*70%", "=ROUNDDOWN(AD{r}*30%,-3)"),
        ("김기도", "=ROUNDDOWN(AA{r}+AB{r}-AC{r},-3)", "=ROUNDDOWN(AD{r}*15%,-3)"),
    ])
    path = tmp_path / "성과상여.xlsx"
    wb.save(path)
    return path


@pytest.fixture()
def db():
    engine = create_engine("sqlite://", future=True, connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine, tables=[BonusPerson.__table__, BonusRate.__table__, BonusShare.__table__])
    session = sessionmaker(bind=engine, future=True)()
    try:
        yield session
    finally:
        session.close()


# ── 읽기 ────────────────────────────────────────────────────────────────

def test_요율_블록은_가장_최근_시트의_것을_쓴다(workbook):
    blocks = excel_seed.read_rate_blocks(excel_seed.open_formulas(workbook))
    assert blocks["강무진"]["sheet"] == "26.08월-주주"
    assert [(b["label"], b["rate"]) for b in blocks["강무진"]["blocks"]] == [
        ("2019년분/2021년 3월~", 40), ("2020년~2021년 2월", 45),
    ]
    assert [(b["label"], b["rate"]) for b in blocks["노승환"]["blocks"]] == [
        ("2025년 7월 7일 접수분~", 40), ("2022년 7월~ 2025년 7월 6일 접수분까지", 30),
    ]
    assert "총괄표 08" not in {v["sheet"] for v in blocks.values()}


def test_지분은_블록_주인에게_귀속하고_수식이_없으면_값으로_읽는다(workbook):
    rows = excel_seed.read_share_rows(excel_seed.open_formulas(workbook), excel_seed.open_values(workbook))
    by_key = {(r["doc_id"], r["person"]): r for r in rows}
    assert by_key[("01-2503-5-0042", "강무진")]["share_pct"] == 25.0        # 수식 *25%
    assert by_key[("01-2503-5-0042", "김형식")]["share_pct"] == 50.0
    assert by_key[("01-2603-1-0155", "안창덕")]["share_pct"] == 90.0        # 값 255,600 / 284,000
    assert by_key[("01-2607-3-0001", "안창덕")]["bc_pct"] == 100.0         # 법카 안100
    assert ("01-2607-3-0002", "김형식") not in by_key                     # F+G = 0
    assert ("01-2606-3-1813", "김형식") not in by_key                     # 혼자 100% 는 저장 안 함
    assert by_key[("01-2603-1-0155", "안창덕")]["note"].startswith("주택도시보증공사")


def test_소속_파라미터는_합계행_수식에서_읽는다(workbook):
    params = excel_seed.read_associate_params(excel_seed.open_formulas(workbook))
    assert params["이영은"] == {"pay_ratio": 0.7, "tax_rate": 0.30, "sheet": "26.08월-평.동"}
    assert params["김기도"] == {"pay_ratio": 1.0, "tax_rate": 0.15, "sheet": "26.08월-평.동"}


def test_사람_목록은_주주와_소속을_합친다(workbook):
    wb = excel_seed.open_formulas(workbook)
    people = excel_seed.read_person_list(excel_seed.read_rate_blocks(wb), excel_seed.read_associate_params(wb))
    by_name = {p["person"]: p for p in people}
    assert by_name["강무진"] == {"person": "강무진", "kind": "SHAREHOLDER", "pay_ratio": 1.0, "tax_rate": 0.30}
    assert by_name["이영은"] == {"person": "이영은", "kind": "ASSOCIATE", "pay_ratio": 0.7, "tax_rate": 0.30}
    # 같은 달에 주주 블록과 소속 합계가 다 있으면 소속이다 (접수일이 구간에 들면 리포트가 주주 행으로 올린다)
    assert by_name["김기도"]["tax_rate"] == 0.15


# ── 저장 ────────────────────────────────────────────────────────────────

def test_미리보기는_아무것도_쓰지_않고_수만_센다(workbook, db):
    report = seed(db, workbook, dry_run=True)
    assert report["dry_run"] is True
    assert report["persons"]["written"] == 7           # 강무진·안창덕·김형식·노승환·황인석 + 이영은·김기도 (황인석은 요율이 없어도 사람이다)
    assert report["rates"]["written"] == 5             # 강무진·안창덕·김형식·노승환·이영은 (황인석은 라벨 미해석)
    assert report["shares"]["written"] == 3            # 0042(2명)·0155·0001 → 감정서 3건
    assert report["unknown_labels"] == [("황인석", "이상한 라벨")]
    assert persons.list_persons(db) == [] and schedule.load_schedule(db) == {}


def test_시드는_쓰고_다시_돌리면_건너뛰며_force_로만_덮어쓴다(workbook, db):
    first = seed(db, workbook, known_names={"강무진", "안창덕", "김형식", "노승환", "이영은", "김기도"})
    assert first["persons"]["written"] == 7 and first["rates"]["written"] == 5 and first["shares"]["written"] == 3
    assert first["unknown_names"] == ["황인석"]
    blocks = schedule.load_schedule(db)["노승환"]
    assert schedule.rate_for(blocks, date(2025, 7, 7)) == (40.0, date(2025, 7, 7), "SCHEDULE")
    assert schedule.rate_for(blocks, date(2025, 7, 6)) == (30.0, date(2022, 7, 1), "SCHEDULE")
    assert schedule.load_schedule(db)["강무진"][0].from_date == date(2019, 1, 1)   # 2019년분
    loaded = shares.load_shares(db, ["01-2503-5-0042", "01-2607-3-0001"])
    assert loaded["01-2503-5-0042"]["강무진"]["share_pct"] == 25.0
    assert loaded["01-2503-5-0042"]["강무진"]["source"] == "SEED"
    assert loaded["01-2607-3-0001"]["안창덕"]["bc_pct"] == 100.0
    assert {p["person"]: p["kind"] for p in persons.list_persons(db)}["이영은"] == "ASSOCIATE"

    again = seed(db, workbook)
    assert (again["persons"]["written"], again["rates"]["written"], again["shares"]["written"]) == (0, 0, 0)
    assert again["persons"]["skipped"] == 7 and again["rates"]["skipped"] == 5 and again["shares"]["skipped"] == 3

    schedule.save_rates(db, "노승환", [{"from_date": None, "to_date": None, "rate": 45}], usr_seq=1)
    forced = seed(db, workbook, force=True)
    assert forced["rates"]["written"] == 5
    assert [b.rate for b in schedule.load_schedule(db)["노승환"]] == [30.0, 40.0]
