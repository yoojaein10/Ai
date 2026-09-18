"""업무실적 보고 — 순수 함수 단위테스트 + 실DB 회귀(2026-05 상반 실파일 대조)."""

from datetime import datetime
from pathlib import Path

import pytest

from app.config import get_settings
from app.services.work_report import (
    _row_to_dict,
    _serialize,
    half_month_range,
    quarter_bungi,
)

# 실파일 후보 — 저장소본을 먼저 본다. 개인 PC 경로만 두면 다른 장비·서버에서
# 영원히 스킵돼 회귀 안전망이 꺼진 채로 지나간다(2026-08-06 실제로 겪음).
TEMPLATE_CANDIDATES = (
    Path(__file__).parents[1] / "docs" / "★26.05월업무실적보고(수수료확인용)-상반.xlsx",
    Path('C:\\Users\\PUBLIC_USER\\Documents\\DHAPPMessenger\\★26.05월업무실적보고(보고용)-상반.xls'),
)
TEMPLATE = next((p for p in TEMPLATE_CANDIDATES if p.exists()), TEMPLATE_CANDIDATES[0])
# 실파일의 수기 추가분(어떤 자동 신호로도 안 나오는 건) — recall 계산에서 제외한다.
MANUAL_DOCS = {
    "01-2604-3-1193", "01-2509-B-0048", "01-2603-1-0104", "01-2603-1-0231",
    "01-2601-5-0003",
}


def test_half_month_range_splits_at_15th():
    assert half_month_range(2026, 5, "상반") == (
        datetime(2026, 5, 1), datetime(2026, 5, 15, 23, 59, 59),
    )
    # 하반은 16일부터 말일까지 (5월=31일)
    assert half_month_range(2026, 5, "하반") == (
        datetime(2026, 5, 16), datetime(2026, 5, 31, 23, 59, 59),
    )


def test_half_month_range_uses_actual_month_end():
    # 2월(윤년 아님)은 28일까지
    assert half_month_range(2026, 2, "하반")[1] == datetime(2026, 2, 28, 23, 59, 59)


def test_quarter_bungi_is_year_plus_quarter():
    assert quarter_bungi(2026, 5) == "20262"   # 5월 → 2분기
    assert quarter_bungi(2026, 1) == "20261"
    assert quarter_bungi(2026, 12) == "20264"


def test_row_to_dict_names_anonymous_address_columns():
    columns = ["ID_NUM", "PCODE"] + [""] * 3  # 18~ 위치가 아니어도 fallback 이름 부여
    row = ("01-2604-3-0001", "82", "a", "b", "c")
    result = _row_to_dict(columns, row)
    assert result["ID_NUM"] == "01-2604-3-0001"
    assert result["PCODE"] == "82"
    # 이름 없는 컬럼은 col{index} 로 (주소 고정위치 18~22가 아니면)
    assert result["col2"] == "a"


def test_serialize_normalizes_types():
    from decimal import Decimal
    assert _serialize(Decimal("1052000")) == 1052000.0
    assert _serialize(datetime(2026, 5, 4)) == "2026-05-04T00:00:00"
    assert _serialize("  01-2604  ") == "01-2604"


# --- 실DB 회귀: 2026-05 상반(본사) 실파일과 대조 ---

def _template_sheet() -> tuple[list[str], dict[str, list]]:
    """실파일 '데이터' 시트 → (헤더 목록, {감정서번호: 행값}). .xls/.xlsx 둘 다 연다.

    보고용(.xls)과 수수료확인용(.xlsx)은 열 배치가 다르다 — 후자에는 요율·격차율·
    의견 열이 끼어 있어 시군구(REG)가 21이 아니라 28이다. 그래서 인덱스를 박지 말고
    헤더 이름으로 찾는다.
    """
    if TEMPLATE.suffix.lower() == ".xlsx":
        openpyxl = pytest.importorskip("openpyxl")
        book = openpyxl.load_workbook(str(TEMPLATE), data_only=True, read_only=True)
        try:
            grid = [list(r) for r in
                    book[book.sheetnames[0]].iter_rows(min_row=1, values_only=True)]
        finally:
            book.close()
    else:
        xlrd = pytest.importorskip("xlrd")
        sheet = xlrd.open_workbook(str(TEMPLATE)).sheet_by_index(0)
        grid = [sheet.row_values(r) for r in range(sheet.nrows)]

    header_at = next(
        i for i, r in enumerate(grid[:5])
        if any("감정서번호" in str(c or "") for c in r)
    )
    header = [str(c or "").replace("\n", " ").strip() for c in grid[header_at]]
    rows = {
        str(r[2]).strip(): ["" if v is None else v for v in r]
        for r in grid[header_at + 1:] if len(r) > 2 and r[2] and str(r[2]).strip()
    }
    return header, rows


def _template_rows() -> dict[str, list]:
    return _template_sheet()[1]


def _template_col(header: list[str], *keywords: str) -> int:
    """헤더에서 keyword 를 모두 포함하는 열의 인덱스."""
    for i, h in enumerate(header):
        if all(k in h for k in keywords):
            return i
    raise AssertionError(f"양식에서 {keywords} 열을 찾지 못했습니다: {header[:16]}")


def _template_docs() -> set[str]:
    """실파일 '데이터' 시트의 감정서번호(3번째 열, 3행부터)."""
    if not TEMPLATE.exists():
        pytest.skip(f"협회 양식 실파일이 없습니다: {TEMPLATE}")
    return set(_template_rows())


@pytest.mark.integration
def test_sales_basis_reproduces_actual_file_recall():
    """매출기준으로 뽑은 감정서가 실제 협회 제출 파일의 자동분(수기 제외)을 재현한다."""
    if not get_settings().is_database_configured:
        pytest.skip("DB 미설정")
    from app.database import get_session_factory

    expected_core = _template_docs() - MANUAL_DOCS
    from app.services.work_report import build_kapa_rows

    session = get_session_factory()()
    try:
        result = build_kapa_rows(
            session, office_id="10", year=2026, month=5, half="상반", basis="매출"
        )
    finally:
        session.close()

    produced = {row["ID_NUM"] for row in result["rows"]}
    recovered = expected_core & produced
    # POC 기준 recall 123/124. 데이터 드리프트 여유로 120 이상이면 통과.
    assert len(recovered) >= 120, (
        f"recall {len(recovered)}/{len(expected_core)}, "
        f"누락 예: {sorted(expected_core - produced)[:5]}"
    )
    assert result["bungi"] == "20262"


# 2026-07 하반 본사 — 담당자가 "실적에 나와야 하는데 안 나온다"고 짚어준 건들.
# 전부 발송은 반월 안에 끝났고 입금·매출계상은 아직 없다(=발송기준으로만 잡힌다).
JULY_SEND_BASIS_DOCS = {
    "01-2607-1-0403",                                               # 3C (공)사용료
    "01-2607-3-2268", "01-2607-3-2269", "01-2607-3-2270",           # 5A HUG(기타담보)
    "01-2607-3-2271", "01-2607-3-2283", "01-2607-3-2292",
    "01-2607-3-2293", "01-2607-3-2330", "01-2607-3-2356",
    "01-2607-3-2371", "01-2607-3-2372", "01-2607-3-2373",
    "01-2607-3-2381", "01-2607-3-2410",
    "01-2607-4-0251", "01-2607-4-0252",                             # 97 HUG(시가참고)
    "01-2607-4-0259", "01-2607-4-0260",
}
# 발송기준 목적이 아니고 입금도 없는 건 — 실적에 들어오면 안 된다.
# 8037은 개별공시지가 검증(12), 나머지 넷은 담보·공매로 전부 SP 목록 밖이다.
# (1065는 5A지만 발주처가 KB자산운용이라 HUG 갈래를 타지 않는다)
JULY_EXCLUDED_DOCS = {
    "01-2607-8-0037",                   # 12 개별공시지가 검증
    "01-2604-3-1065",                   # 5A 기타담보 · HUG 아님
    "01-2607-2-0088",                   # 9F 공매(NPL)
    "01-2607-3-2322", "01-2607-3-2325",  # 52 제1금융권담보
}


@pytest.mark.integration
def test_sales_basis_includes_hug_and_compensation_by_send_date():
    """발송기준 목적(HUG·보상·공공)은 입금 전이라도 반월 실적에 들어간다."""
    if not get_settings().is_database_configured:
        pytest.skip("DB 미설정")
    from app.database import get_session_factory
    from app.services.work_report import build_kapa_rows

    session = get_session_factory()()
    try:
        result = build_kapa_rows(
            session, office_id="10", year=2026, month=7, half="하반", basis="매출"
        )
    finally:
        session.close()

    produced = {row["ID_NUM"] for row in result["rows"]}
    missing = JULY_SEND_BASIS_DOCS - produced
    assert not missing, f"발송기준인데 빠진 감정서: {sorted(missing)}"
    leaked = JULY_EXCLUDED_DOCS & produced
    assert not leaked, f"입금도 발송기준도 아닌데 들어온 감정서: {sorted(leaked)}"


def test_excel_builder_maps_columns_by_position():
    from openpyxl import load_workbook
    from io import BytesIO
    from app.services.work_report_excel import build_work_report_xlsx

    result = {
        "bungi": "20262", "month": 5, "appcode": "300611",
        "rows": [{
            "ID_NUM": "01-2604-1-0254", "QUANO": "2405", "GAMMAN": "고세욱",
            "PCODE": "26", "YCODE": "00", "GNAME": "보상감정", "CUST": "보성군수",
            "CUSTCODE": "1", "CONSULTDATE": "2026-05-22T00:00:00",
            "IN_DATE": "2026-04-10T00:00:00", "OUTDATE": "2026-05-07T00:00:00",
            "GAMGA": 3037586440.0, "FEE": 2266317.0, "SUSU": 2266317.0,
            "REG": "46780", "EUB": "25321", "SAN": "1", "BUN1": "0875", "BUN2": "0040",
            "CATEGORY1": "1", "CNT1": 4, "PRICE1": 2883270000.0,
            "CATEGORY2": "2", "CNT2": 14, "PRICE2": 154316440.0,
        }],
    }
    book = load_workbook(BytesIO(build_work_report_xlsx(result)))
    sheet = book["데이터"]
    assert sheet["A3"].value == "20262"      # 분기
    assert sheet["B3"].value == 5            # 월
    assert sheet["C3"].value == "01-2604-1-0254"
    assert sheet["D3"].value == "2405"       # 자격번호(문자, 앞0 보존)
    assert sheet["H3"].value == "26"         # 평가목적 협회코드
    assert sheet["Q3"].value == "2026-04-10"  # 접수일 (날짜만)
    assert sheet["S3"].value == 3037586440.0  # 감정평가액(숫자)
    assert sheet["V3"].value == "11380" or sheet["V3"].value == "46780"  # 시군구
    # 물건구분1: 협회코드/물건수/평가액 = AA/AB/AC
    assert sheet["AA3"].value == "1" and sheet["AB3"].value == 4
    assert sheet["AC3"].value == 2883270000.0
    # 협회코드 참조시트가 그대로 살아있어야 한다
    assert "협회코드" in book.sheetnames


def test_send_basis_follows_hug_branch_of_sp():
    """HUG(주택도시보증공사) 건은 SP가 5A·97·46·33을 발송기준으로 본다."""
    from app.services.work_report import _send_basis_purposes_sql

    sql = _send_basis_purposes_sql()
    assert "주택도시보증공사" in sql
    assert "2026-02-27" in sql          # SP가 HUG 갈래를 적용하는 발송일 하한
    for code in ("5A", "97", "46", "33"):
        assert f"'{code}'" in sql, code


def test_send_basis_follows_general_branch_of_sp():
    """HUG가 아닌 건은 SP의 범위(21~27·31~3C·62~67)와 목록을 그대로 쓴다."""
    from app.services.work_report import _send_basis_purposes_sql

    sql = _send_basis_purposes_sql()
    for low, high in (("21", "27"), ("31", "3C"), ("62", "67")):
        assert f"BETWEEN '{low}' AND '{high}'" in sql, (low, high)
    assert "'39', '3A'" in sql          # 31~3C 안에서 SP가 빼는 두 코드
    for code in ("72", "75", "76", "77", "92", "93", "95", "9C", "9E", "9H", "9I"):
        assert f"'{code}'" in sql, code


def test_send_basis_keeps_court_auction_exception():
    """61(법원경매)은 SP 범위(62~67)에서 빠져 있으나 실파일에는 있어 예외로 넣는다."""
    from app.services.work_report import SEND_BASIS_EXTRA_CODES, _send_basis_purposes_sql

    assert "61" in SEND_BASIS_EXTRA_CODES
    assert "'61'" in _send_basis_purposes_sql()


def test_sigacham_is_send_basis_only_for_hug():
    """시가참고(97)는 SP의 HUG 갈래에만 있다 — HUG가 아닌 97은 발송으로 안 들어온다."""
    from app.services.work_report import _send_basis_purposes_sql

    sql = _send_basis_purposes_sql()
    hug_clause = sql[:sql.index("NOT (m.Production")]
    assert "'97'" in hug_clause          # HUG 갈래 안에는 있고
    assert sql.count("'97'") == 1        # 그 밖에는 어디에도 없다


def test_send_branch_ignores_sales_posting():
    """발송 브랜치는 발송기준 목적만 본다 — 매출계상(#MAE)은 판단 근거가 아니다.

    담당자가 짚어준 4건(01-2604-3-1065 등)이 입금도 없이 '7월에 전표가 세워졌다'는
    이유만으로 실적에 들어오던 원인이라 갈래째 뺐다.
    """
    from app.services.work_report import _select_sales_send_only_sql, _select_sales_sql

    assert "#MAE" not in _select_sales_sql()
    assert "#MAE" not in _select_sales_send_only_sql()


def test_sales_send_only_sql_has_no_deposit_branch():
    """수수료 계상 없음 판정용 발송전용 SQL — 입금(#IN) 브랜치 없이 발송 조건만."""
    from app.services.work_report import _select_sales_send_only_sql, _select_sales_sql

    sql = _select_sales_send_only_sql()
    assert "#IN" not in sql
    assert "m.SendDate BETWEEN ? AND ?" in sql
    # 바인드 순서 계약: 공통 3개(appcode/bungi/mon) + 발송창 2개(start/end)
    assert sql.count("?") == 5
    # 발송 브랜치 문구는 기존 선택 SQL과 동일해야 한다(한쪽만 고치는 회귀 방지)
    branch = sql[sql.index("m.SendDate"):]
    assert branch in _select_sales_sql()


def test_docs_with_fee_voucher_chunks_and_strips():
    """전표캐시에서 감정수수료(4010001) 계상 감정서를 청크로 나눠 조회한다."""
    from app.services.work_report import _docs_with_fee_voucher

    class FakeDb:
        def __init__(self):
            self.calls = []

        def execute(self, statement, params):
            self.calls.append(params)
            # 넘어온 것 중 '-ok'로 끝나는 감정서만 계상 있음으로 응답 (공백 포함)
            return [(f" {doc} ",) for doc in params.values() if doc.endswith("-ok")]

    fake = FakeDb()
    docs = [f"01-2607-3-{i:04d}-ok" for i in range(600)] + ["01-2604-3-1309"]
    found = _docs_with_fee_voucher(fake, docs)
    assert len(fake.calls) == 2  # 500개 청크 2번
    assert "01-2607-3-0000-ok" in found and "01-2607-3-0599-ok" in found
    assert "01-2604-3-1309" not in found
    assert _docs_with_fee_voucher(fake, []) == set()


def test_cell_value_normalizes_kinds():
    from app.services.work_report_excel import _cell_value
    assert _cell_value("date", "2026-05-04T00:00:00") == "2026-05-04"
    assert _cell_value("num", "1052000") == 1052000.0
    assert _cell_value("int", 4.0) == 4
    assert _cell_value("text", "  0973 ") == "0973"
    assert _cell_value("text", "") is None
    assert _cell_value("num", None) is None


def test_compute_ak_matches_sample_formula():
    from datetime import date
    from app.services.kapa_submit import compute_ak
    # upmu 샘플: 2022-09-26 → 2022*9*26*261 = 123491628
    assert compute_ak(date(2022, 9, 26)) == "123491628"
    assert compute_ak(date(2026, 5, 4)) == str(2026 * 5 * 4 * 261)


def test_build_datas_field_order_and_types():
    import json
    from app.services.kapa_submit import build_datas, _DATAS_FIELDS

    rows = [{
        "ID_NUM": "01-2604-1-0254", "QUANO": "2405", "GNAME": "보상감정",
        "IN_DATE": "2026-04-10T00:00:00", "OUTDATE": "2026-05-07T00:00:00",
        "CONSULTDATE": "2026-05-22T00:00:00", "GAMGA": 3037586440.0, "CATEGORY1": "1",
        "GAMMAN": "고세욱",  # 성명은 DATAS에 포함되면 안 됨
    }]
    datas = build_datas(rows)
    obj = json.loads(datas)[0]
    assert list(obj.keys()) == _DATAS_FIELDS       # 정확한 필드 순서
    assert "GAMMAN" not in obj                     # 성명 제외
    assert obj["INDATE"] == "2026-04-10"           # IN_DATE→INDATE, 날짜만
    assert obj["GAMGA"] == "3037586440.0"          # 모든 값 문자열
    assert obj["QUANO"] == "2405"
    assert " " not in datas and "\n" not in datas  # 공백/개행 없음


def test_fetch_registered_uses_query_constant_and_parses_ids(monkeypatch):
    """조회는 AK상수 137(전송 261과 다름), DATA에서 ID_NUM 집합을 뽑는다."""
    from datetime import date
    import app.services.kapa_submit as ks

    captured = {}

    class FakeResponse:
        status_code = 200
        def json(self):
            return {"RESULT": "0", "MSG": "", "DATA": [
                {"ID_NUM": "01-2604-1-0254 "}, {"ID_NUM": ""}, {"no_id": 1},
            ]}

    def fake_post(url, data=None, timeout=None):
        captured.update(url=url, ak=data["AK"])
        return FakeResponse()

    monkeypatch.setattr(ks.requests, "post", fake_post)
    settings = ks.get_settings()
    monkeypatch.setattr(type(settings), "is_kapa_configured", property(lambda self: True))
    ids = ks.fetch_registered_doc_ids("300611", "20263", 7, today=date(2026, 7, 27))
    assert ids == {"01-2604-1-0254"}
    assert captured["url"].endswith("/rest/LAWREP/")
    assert captured["ak"] == str(2026 * 7 * 27 * 137)   # 조회 상수
    # 전송 기본값은 여전히 261
    assert ks.compute_ak(date(2026, 7, 27)) == str(2026 * 7 * 27 * 261)


def test_fetch_registered_raises_on_error_result(monkeypatch):
    import app.services.kapa_submit as ks

    class FakeResponse:
        status_code = 200
        def json(self):
            return {"RESULT": "40", "MSG": "AK Validation Error"}

    monkeypatch.setattr(ks.requests, "post", lambda *a, **k: FakeResponse())
    settings = ks.get_settings()
    monkeypatch.setattr(type(settings), "is_kapa_configured", property(lambda self: True))
    with pytest.raises(ks.KapaQueryError):
        ks.fetch_registered_doc_ids("300611", "20263", 7)


def test_submit_dry_run_never_hits_network_and_hides_credentials(monkeypatch):
    import app.services.kapa_submit as ks
    monkeypatch.setattr(ks.requests, "post", lambda *a, **k: pytest.fail("dry-run이 네트워크 호출"))
    settings = ks.get_settings()
    monkeypatch.setattr(type(settings), "is_kapa_configured", property(lambda self: True))
    result = {"appcode": "300611", "bungi": "20262", "month": 5,
              "rows": [{"ID_NUM": "01-2604-1-0254"}]}
    out = ks.submit_work_report(result, dry_run=True)
    assert out["dry_run"] is True and out["count"] == 1
    # 자격증명이 응답에 새지 않는다
    assert "PASSWD" not in str(out) and "USERID" not in str(out)


@pytest.mark.integration
def test_mapping_columns_match_template_cells():
    """매핑 재현 값이 양식 셀과 일치(핵심 컬럼). 원본 데이터 드리프트분만 예외."""
    if not get_settings().is_database_configured:
        pytest.skip("DB 미설정")
    if not TEMPLATE.exists():
        pytest.skip(f"협회 양식 실파일이 없습니다: {TEMPLATE}")
    from app.database import get_session_factory
    from app.services.work_report import build_kapa_rows

    header, tpl = _template_sheet()
    session = get_session_factory()()
    try:
        result = build_kapa_rows(
            session, office_id="10", year=2026, month=5, half="상반", basis="매출"
        )
    finally:
        session.close()
    mine = {row["ID_NUM"]: row for row in result["rows"]}

    # (양식 헤더 키워드, 매핑 키) — 협회코드/자격번호는 100% 일치해야 한다.
    # 인덱스를 박으면 보고용/수수료확인용 레이아웃 차이에 그대로 깨진다.
    checks = [
        (_template_col(header, "담당평가사1", "자격번호"), "QUANO"),
        (_template_col(header, "협회코드", "평가목적"), "PCODE"),
        (_template_col(header, "평가구분", "코드"), "YCODE"),
        (_template_col(header, "협회코드", "의뢰처"), "CUSTCODE"),
        (_template_col(header, "시군구"), "REG"),
    ]
    shared = set(tpl) & set(mine)
    assert len(shared) >= 120
    for col_index, key in checks:
        mismatched = [
            doc for doc in shared
            if str(tpl[doc][col_index]).strip().lstrip("0")
            != str(mine[doc].get(key, "")).strip().lstrip("0")
        ]
        # 드리프트 여유로 5건 이하 허용
        assert len(mismatched) <= 5, f"{key} 불일치 {len(mismatched)}건: {mismatched[:5]}"
