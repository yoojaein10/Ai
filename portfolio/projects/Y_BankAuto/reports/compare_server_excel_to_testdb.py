import configparser
import datetime as dt
import json
import math
import re
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pyodbc


BASE = Path(r"D:\AI\Claude\Y_BankAuto")
REPORT_DIR = BASE / "reports"
REPORT_DIR.mkdir(parents=True, exist_ok=True)

SOURCE_XLSX = next(BASE.glob("*db.xlsx"))
SETTINGS = Path(r"C:\Bank24Extractor\settings.ini")

QUERY_COLUMNS = [
    "Office", "ReceiptDate", "RequestDate", "Priority", "WorkInfo", "Purpose",
    "Category", "Inventory", "CustID", "CustName", "Production", "CustCharge",
    "CustPhone", "CustDocID", "Debtor", "DebtrPhone", "Title", "Reg", "Eub",
    "SAN", "ADDR", "BUN1", "BUN2", "Building", "Dong", "Floor", "Ho", "hoetc",
]

PHONE_COLUMNS = {"CustPhone", "DebtrPhone"}
NUMERIC_TEXT_COLUMNS = {
    "Office", "Priority", "WorkInfo", "Purpose", "Category", "Inventory",
    "CustID", "Reg", "Eub", "SAN", "BUN1", "BUN2",
}
TEXT_SPACE_COLUMNS = {
    "CustName", "Production", "CustCharge", "Debtor", "Title", "ADDR",
    "Building", "Dong", "Floor", "Ho", "hoetc",
}
DATE_COLUMNS = {"ReceiptDate", "RequestDate"}


def db_conn_str():
    cfg = configparser.ConfigParser()
    cfg.read(SETTINGS, encoding="utf-8")
    db = cfg["database"]
    driver = db.get("driver", "ODBC Driver 17 for SQL Server")
    port = db.get("port", "1433")
    return (
        f"DRIVER={{{driver}}};SERVER={db.get('server')},{port};DATABASE=apworksdw;"
        f"UID={db.get('username')};PWD={db.get('password')};TrustServerCertificate=yes;"
    )


def raw_value(value):
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value).strip()


def is_nullish(text):
    return text.strip().upper() in {"", "NULL", "NONE", "NAN"}


def normalize_for_compare(column, value):
    text = raw_value(value)
    if is_nullish(text):
        return ""
    if column in DATE_COLUMNS:
        m = re.match(r"^(\d{4}-\d{2}-\d{2})", text)
        return m.group(1) if m else text
    if column in PHONE_COLUMNS:
        return re.sub(r"\D+", "", text)
    if column in NUMERIC_TEXT_COLUMNS:
        digits = re.sub(r"\D+", "", text)
        if digits:
            return str(int(digits))
    if column in TEXT_SPACE_COLUMNS:
        return re.sub(r"\s+", " ", text).strip()
    return re.sub(r"\s+", " ", text).strip()


def custdoc_from_excel(value):
    raw = raw_value(value)
    if is_nullish(raw):
        return "", "", "empty"
    raw = raw.strip()
    if raw.startswith("'"):
        raw = raw[1:].strip()
        return raw, raw, "exact"
    # Excel numeric cells lose precision for 16+ digit IDs. Preserve what exists,
    # and keep a prefix token for best-effort unique matching.
    if isinstance(value, (int, float)) or re.fullmatch(r"\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?", raw):
        try:
            dec = Decimal(str(value))
            fixed = format(dec.quantize(Decimal("1")), "f")
        except (InvalidOperation, ValueError):
            fixed = raw
        fixed = re.sub(r"\D+", "", fixed)
        if len(fixed) >= 16:
            return fixed, fixed[:13], "excel_numeric_precision"
        return fixed, fixed, "exact"
    cleaned = re.sub(r"\.0$", "", raw)
    cleaned = cleaned.strip()
    return cleaned, cleaned, "exact"


def read_server_excel():
    wb = openpyxl.load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    rows = []
    notes = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        try:
            header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        except StopIteration:
            continue
        headers = [raw_value(h) for h in header_row]
        if not headers or "CustDocID" not in headers:
            continue
        index = {h: i for i, h in enumerate(headers)}
        for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(not is_nullish(raw_value(v)) for v in row):
                continue
            rec = {}
            for col in QUERY_COLUMNS:
                excel_col = "ho" if col == "Ho" and "ho" in index else col
                rec[col] = row[index[excel_col]] if excel_col in index and index[excel_col] < len(row) else ""
            cust_raw, match_key, match_mode = custdoc_from_excel(rec["CustDocID"])
            rec["CustDocID_raw"] = cust_raw
            rec["CustDocID_match_key"] = match_key
            rec["CustDocID_match_mode"] = match_mode
            rec["_sheet"] = sheet_name
            rec["_row"] = row_no
            rows.append(rec)
            if match_mode == "excel_numeric_precision":
                notes.append({
                    "Sheet": sheet_name,
                    "Row": row_no,
                    "ExcelCustDocID": raw_value(rec["CustDocID"]),
                    "RecoveredDigits": cust_raw,
                    "MatchPrefix": match_key,
                    "Note": "Excel numeric precision loss; matched only if prefix is unique in test DB.",
                })
    return rows, notes


def fetch_test_db_rows():
    sql = """
    SELECT Office
          , ReceiptDate
          , RequestDate
          , Priority
          , WorkInfo
          , Purpose
          , Category
          , Inventory
          , CustID
          , CustName
          , Production
          , CustCharge
          , CustPhone
          , CustDocID
          , Debtor
          , DebtrPhone
          , Title
          , b.Reg
          , b.Eub
          , b.SAN
          , b.ADDR
          , b.BUN1
          , b.BUN2
          , b.Building
          , b.Dong
          , b.Floor
          , b.ho
          , b.hoetc
    FROM Apw_Master a
        INNER JOIN APW_Inventory b
            ON a.MasterID = b.MasterID
           AND a.DocID = b.DocID
    WHERE 1=1
      AND CONVERT(varchar(20), Receiptdate, 23) >= '2026-06-05'
    """
    with pyodbc.connect(db_conn_str(), timeout=15) as conn:
        cur = conn.cursor()
        cur.execute(sql)
        out = []
        db_cols = [d[0] for d in cur.description]
        col_map = {c.lower(): c for c in db_cols}
        for row in cur.fetchall():
            rec = {}
            for col in QUERY_COLUMNS:
                db_col = col_map.get(col.lower())
                rec[col] = getattr(row, db_col) if db_col else ""
            out.append(rec)
        return out


def inventory_signature(row):
    return tuple(normalize_for_compare(c, row.get(c)) for c in ["Reg", "Eub", "SAN", "ADDR", "BUN1", "BUN2", "Building", "Dong", "Floor", "Ho", "hoetc"])


def pair_rows(server_rows, test_rows):
    pairs = []
    used = set()
    for s in server_rows:
        sig = inventory_signature(s)
        exact = None
        for i, t in enumerate(test_rows):
            if i in used:
                continue
            if inventory_signature(t) == sig:
                exact = i
                break
        if exact is not None:
            used.add(exact)
            pairs.append((s, test_rows[exact]))
            continue
        best_i = None
        best_score = -1
        for i, t in enumerate(test_rows):
            if i in used:
                continue
            score = sum(
                normalize_for_compare(c, s.get(c)) == normalize_for_compare(c, t.get(c))
                for c in ["Reg", "Eub", "SAN", "ADDR", "BUN1", "BUN2", "Building", "Dong", "Floor", "Ho", "hoetc"]
            )
            if score > best_score:
                best_i, best_score = i, score
        if best_i is not None:
            used.add(best_i)
            pairs.append((s, test_rows[best_i]))
        else:
            pairs.append((s, None))
    for i, t in enumerate(test_rows):
        if i not in used:
            pairs.append((None, t))
    return pairs


server_rows, precision_notes = read_server_excel()
test_rows = fetch_test_db_rows()

test_by_doc = defaultdict(list)
test_prefix = defaultdict(set)
test_numeric_key = defaultdict(set)
for row in test_rows:
    doc = normalize_for_compare("CustDocID", row.get("CustDocID"))
    test_by_doc[doc].append(row)
    if len(doc) >= 13:
        test_prefix[doc[:13]].add(doc)
    if doc.isdigit():
        test_numeric_key[doc.lstrip("0") or "0"].add(doc)

server_by_doc = defaultdict(list)
unmatchable_excel = []
for row in server_rows:
    mode = row["CustDocID_match_mode"]
    key = row["CustDocID_match_key"]
    doc = row["CustDocID_raw"]
    if mode == "excel_numeric_precision":
        candidates = sorted(test_prefix.get(key, set()))
        if len(candidates) == 1:
            doc = candidates[0]
        else:
            unmatchable_excel.append({
                "Sheet": row["_sheet"],
                "Row": row["_row"],
                "ExcelCustDocID": raw_value(row["CustDocID"]),
                "RecoveredDigits": row["CustDocID_raw"],
                "MatchPrefix": key,
                "CandidateCount": len(candidates),
                "Candidates": ", ".join(candidates[:10]),
                "Reason": "Long numeric CustDocID in Excel is not exact and prefix is not uniquely matchable.",
            })
            continue
    elif doc not in test_by_doc and doc.isdigit():
        # Excel numeric formatting removes leading zeroes for IDs like 0010260313.
        candidates = sorted(test_numeric_key.get(doc.lstrip("0") or "0", set()))
        if len(candidates) == 1:
            doc = candidates[0]
    if doc:
        server_by_doc[doc].append(row)

common_docs = sorted(set(server_by_doc) & set(test_by_doc))
server_only_docs = sorted(set(server_by_doc) - set(test_by_doc))
test_only_docs_skipped = sorted(set(test_by_doc) - set(server_by_doc))

summary_by_doc = []
diff_rows = []
field_counts = Counter()

for doc in common_docs:
    s_rows = server_by_doc[doc]
    t_rows = test_by_doc[doc]
    pairs = pair_rows(s_rows, t_rows)
    doc_diff_count = 0
    for idx, (s, t) in enumerate(pairs, start=1):
        if s is None or t is None:
            doc_diff_count += 1
            diff_rows.append({
                "CustDocID": doc,
                "PairIndex": idx,
                "Field": "RowPresence",
                "ServerValue": "MISSING" if s is None else "EXISTS",
                "TestValue": "MISSING" if t is None else "EXISTS",
                "ServerSheet": "" if s is None else s["_sheet"],
                "ServerRow": "" if s is None else s["_row"],
            })
            field_counts["RowPresence"] += 1
            continue
        for col in QUERY_COLUMNS:
            if col == "CustDocID":
                continue
            sv = normalize_for_compare(col, s.get(col))
            tv = normalize_for_compare(col, t.get(col))
            if sv != tv:
                doc_diff_count += 1
                field_counts[col] += 1
                diff_rows.append({
                    "CustDocID": doc,
                    "PairIndex": idx,
                    "Field": col,
                    "ServerValue": raw_value(s.get(col)),
                    "TestValue": raw_value(t.get(col)),
                    "ServerSheet": s["_sheet"],
                    "ServerRow": s["_row"],
                })
    summary_by_doc.append({
        "CustDocID": doc,
        "ServerRows": len(s_rows),
        "TestRows": len(t_rows),
        "ComparedPairs": len(pairs),
        "DiffCount": doc_diff_count,
        "Status": "MATCH" if doc_diff_count == 0 else "DIFF",
    })

payload = {
    "created_at": dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    "source_excel": str(SOURCE_XLSX),
    "test_query": "APW_Master INNER JOIN APW_Inventory ON MasterID+DocID, ReceiptDate >= 2026-06-05",
    "summary_metrics": [
        {"Metric": "ServerExcelRows", "Value": len(server_rows)},
        {"Metric": "ServerExcelMatchableCustDocID", "Value": len(server_by_doc)},
        {"Metric": "TestDBRows", "Value": len(test_rows)},
        {"Metric": "CommonCustDocIDCompared", "Value": len(common_docs)},
        {"Metric": "ServerCustDocIDNotInTestDB", "Value": len(server_only_docs)},
        {"Metric": "TestCustDocIDSkippedNotInExcel", "Value": len(test_only_docs_skipped)},
        {"Metric": "UnmatchableExcelNumericCustDocIDRows", "Value": len(unmatchable_excel)},
        {"Metric": "DiffCustDocIDCount", "Value": len([r for r in summary_by_doc if r["Status"] == "DIFF"])},
        {"Metric": "DiffCellOrRowCount", "Value": len(diff_rows)},
    ],
    "summary_by_doc": summary_by_doc,
    "diff_rows": diff_rows,
    "field_counts": [{"Field": k, "DiffCount": v} for k, v in field_counts.most_common()],
    "server_only_docs": [{"CustDocID": x, "Reason": "In server Excel but not in test DB query result"} for x in server_only_docs],
    "test_only_docs_skipped": [{"CustDocID": x, "Reason": "In test DB query result but not in server Excel; skipped by request"} for x in test_only_docs_skipped],
    "excel_precision_notes": precision_notes,
    "unmatchable_excel": unmatchable_excel,
}

out = REPORT_DIR / f"server_excel_vs_testdb_{dt.datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
print(out)
for item in payload["summary_metrics"]:
    print(f"{item['Metric']}: {item['Value']}")
print("TopFields:", payload["field_counts"][:12])
